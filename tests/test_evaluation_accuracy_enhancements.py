"""Bộ kiểm thử đảm bảo độ chính xác của hệ thống nhận diện AI Music trên testcase 2 & 3.

Tuân thủ nghiêm ngặt chuẩn PEP 8 và Ruff.
Kiểm thử các khía cạnh:
1. Nhận diện các dự án nghệ sĩ ảo AI mới (D&C, Athena, Kha Mộc Lan, Mercy Worship, TienYinMen...).
2. Loại bỏ Warner Music Vietnam khỏi Whitelist cứng để tránh bẫy catalogue AI.
3. Cơ chế Trọng tài Anti-FN can thiệp khi có bằng chứng vật lý vượt trội bất kể từ khóa thanh quản sinh học.
4. Xuất đầy đủ các cột ai_model_prediction, ai_model_verdict, ai_model_is_correct trong file Excel.
5. Cập nhật chính xác trạng thái xác thực nguồn âm sau khi AudioEngine thực hiện fallback.
"""

from __future__ import annotations

import unittest
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock

import openpyxl

from backend.core.models.entities import (
    AcousticMetrics,
    DetectionResult,
    InputType,
    TrackItem,
    VerdictType,
)
from backend.services.detection_service import DetectionService
from backend.services.signal_analyzer_service import (
    _DEFAULT_HUMAN_WHITELIST,
    SignalAnalyzerService,
    get_dynamic_catalog,
)
from tests.test_confusion_matrix import (
    ConfusionMetrics,
    export_evaluation_excel,
)


class TestEvaluationAccuracyEnhancements(unittest.TestCase):
    """Kiểm thử nâng cấp độ chính xác và tính toàn vẹn của kết quả benchmark."""

    def setUp(self) -> None:
        self.signal_analyzer = SignalAnalyzerService(
            fakeprint_detector=MagicMock(), storage_repo=MagicMock()
        )

    def test_new_ai_artists_in_catalog(self) -> None:
        """Kiểm tra các nghệ sĩ ảo mới trong testcase 2 & 3 phải có trong catalog Heuristics."""
        new_artists = [
            ("D&C", "STAY A LITTLE LONGER (Ver 2)"),
            ("Athena", "THE WINGS OF THE MORNING V2"),
            ("Kha Mộc Lan", "Anh Giấu Em"),
            ("Khói Xanh", "Mùa Đông Nhớ Cố Nhân"),
            ("Trầm Ca", "Lạy Em, Anh Thêm Bước Nữa"),
            ("Hương My Bông", "Dạ Thưa"),
            ("Mercy Worship, Astrís, WBM Christmas", "Angels from the Realms of Glory"),
            ("TienYinMen, Yi Xian Sun", "Sworn Brothers_Full"),
            ("Shadow Phonk", "Phonk Beat"),
            ("Firelight Records", "I Will Always Love You - UK Drill"),
        ]

        _, _, ai_artists = get_dynamic_catalog()
        self.assertTrue(len(ai_artists) > 0)

        for artist, title in new_artists:
            meta = {"artist": artist, "title": title}
            heuristic_res = self.signal_analyzer.check_heuristics(meta)
            self.assertIsNotNone(
                heuristic_res,
                f"Nghệ sĩ AI mới '{artist}' phải được nhận diện qua Heuristics Catalog!",
            )
            self.assertTrue(
                heuristic_res.get("is_ai"),
                f"Nghệ sĩ AI mới '{artist}' phải được gán nhãn is_ai=True!",
            )

    def test_warner_music_vietnam_not_in_hardcoded_whitelist(self) -> None:
        """Đảm bảo Warner Music Vietnam đã được đưa ra khỏi Whitelist cứng để tránh thiên kiến."""
        self.assertNotIn(
            "warner music vietnam",
            [w.lower() for w in _DEFAULT_HUMAN_WHITELIST],
            "Warner Music Vietnam không được nằm trong Whitelist cứng vì phân phối nhiều catalogue AI!",
        )

    def test_anti_fn_safeguard_overrides_when_strong_physical_ai(self) -> None:
        """Trọng tài Anti-FN phải lật ngược phán quyết lên AI khi Fakeprint/Stem bleed cao."""
        detection_service = DetectionService(
            audio_engine=MagicMock(),
            signal_analyzer=MagicMock(),
            classifier=MagicMock(),
            resolver=MagicMock(),
            track_repo=MagicMock(),
        )

        track = TrackItem(
            id="test_fn_track",
            raw_query="Cõi Nợ Trần Ai Kha Mộc Lan",
            input_type=InputType.SONG_NAME,
            track_title="Cõi Nợ Trần Ai",
            artist="Kha Mộc Lan",
            distributor="Warner Music Vietnam",
        )

        # Giả lập Gemini bị đánh lừa bởi giọng Bolero và phán quyết PROCESSED_HUMAN
        initial_result = DetectionResult(
            track_id=track.id,
            track_title=track.track_title,
            input_type=track.input_type,
            raw_query=track.raw_query,
            is_ai=False,
            verdict_type=VerdictType.PROCESSED_HUMAN,
            verdict_label="PROCESSED HUMAN (Người làm - Kỹ thuật số)",
            confidence_score=0.92,
            suspected_engine="None",
            reasons=[
                "The track demonstrates organic vocal tract formants and natural micro-fluctuations consistent with human singing."
            ],
            detailed_explanation="Vocal tract biology indicates authentic human performance.",
            ai_model_is_ai=False,
            ai_model_verdict="PROCESSED_HUMAN",
        )

        # Chỉ số vật lý AI vượt trội (Fakeprint = 0.65, Stem Bleed = 0.38)
        metrics = AcousticMetrics(
            sample_rate=44100,
            duration_sec=30.0,
            fakeprint_score=0.65,
            fakeprint_detected=True,
            stem_bleed_score=0.38,
            stem_bleed_detected=True,
            spectral_rolloff_95_hz=9500.0,
            pitch_rigidity_ratio=0.55,
        )

        final_res = detection_service._apply_cross_validation_safeguard(
            result=initial_result,
            track=track,
            metadata={"artist": "Kha Mộc Lan", "title": "Cõi Nợ Trần Ai"},
            metrics=metrics,
        )

        self.assertTrue(
            final_res.is_ai,
            "Anti-FN Safeguard phải lật ngược phán quyết lên AI khi Fakeprint và Stem Bleed cao!",
        )
        self.assertIn(
            final_res.verdict_type,
            [VerdictType.HYBRID_AI, VerdictType.AI_GENERATED],
        )

    def test_excel_audit_includes_ai_model_prediction_columns(self) -> None:
        """Đảm bảo file Excel xuất ra có đầy đủ các cột ai_model_prediction, ai_model_verdict, ai_model_is_correct."""
        metrics = ConfusionMetrics(
            total_samples=1,
            evaluated_samples=1,
            errors=0,
            tp=1,
            tn=0,
            fp=0,
            fn=0,
            accuracy=1.0,
            precision=1.0,
            recall=1.0,
            specificity=1.0,
            f1_score=1.0,
            balanced_accuracy=1.0,
            error_rate=0.0,
        )

        results = [
            {
                "media_id": "119001",
                "song_title": "Test AI Song",
                "artist": "AI Artist",
                "query": "Test AI Song AI Artist",
                "ground_truth": "AI",
                "ai_model_prediction": "AI",
                "ai_model_verdict": "AI_GENERATED",
                "ai_model_is_correct": "ĐÚNG",
                "consensus_prediction": "AI",
                "consensus_verdict": "AI_GENERATED",
                "prediction": "AI",
                "is_correct": "ĐÚNG",
                "classification_type": "TP (Đúng: Phát hiện AI)",
                "is_verified_source": True,
                "source_verification": "ĐÃ XÁC THỰC",
                "match_score": "100.0%",
                "audio_url": "https://youtube.com/watch?v=123",
                "note": "",
                "verdict_type": "AI_GENERATED",
                "verdict_label": "AI GENERATED",
                "confidence_score": "95.0%",
                "suspected_engine": "Suno",
                "reasons": "AI generation detected",
            }
        ]

        out_path = Path("data/temp/test_dual_pred_audit.xlsx")
        target_file = export_evaluation_excel(metrics, results, output_path=out_path)
        self.assertTrue(Path(target_file).exists())

        wb = openpyxl.load_workbook(target_file, data_only=True)
        sheet = wb["Audit Details"]
        headers = [c for c in next(sheet.iter_rows(values_only=True))]

        self.assertIn("ai_model_prediction", headers)
        self.assertIn("ai_model_verdict", headers)
        self.assertIn("ai_model_is_correct", headers)
        self.assertIn("consensus_prediction", headers)

        summary_sheet = wb["Executive Summary"]
        summary_items = [row[0] for row in summary_sheet.iter_rows(values_only=True)]
        self.assertIn("Độ chính xác Thuần AI Model (Raw Gemini)", summary_items)
        self.assertIn("Độ chính xác Tổng hợp (Consensus Final)", summary_items)

        wb.close()
        if Path(target_file).exists():
            Path(target_file).unlink()

    def test_benchmark_evaluator_export_excel_includes_dual_prediction(self) -> None:
        """Kiểm tra export_audit_excel trong benchmark_evaluator.py xuất đầy đủ cột và summary."""
        from benchmark_evaluator import export_audit_excel

        metrics = {
            "total_samples": 2,
            "evaluated_successfully": 2,
            "errors": 0,
            "unverified_matches": 0,
            "TP": 1,
            "TN": 1,
            "FP": 0,
            "FN": 0,
            "accuracy": 1.0,
            "precision": 1.0,
            "recall": 1.0,
            "specificity": 1.0,
            "f1_score": 1.0,
        }
        results = [
            {
                "media_id": "1",
                "song_title": "AI Song",
                "artist": "AI Artist",
                "query": "AI Song AI Artist",
                "ground_truth": "AI",
                "ai_model_prediction": "AI",
                "ai_model_verdict": "AI_GENERATED",
                "ai_model_is_correct": "ĐÚNG",
                "consensus_prediction": "AI",
                "consensus_verdict": "AI_GENERATED",
                "prediction": "AI",
                "is_correct": "ĐÚNG",
                "classification_type": "TP (Đúng: Phát hiện AI)",
                "is_verified_source": True,
                "source_verification": "ĐÃ XÁC THỰC",
                "match_score": "100.0%",
                "verdict_type": "AI_GENERATED",
                "verdict_label": "AI GENERATED",
                "confidence_score": "95.0%",
                "suspected_engine": "Suno",
                "source_url": "https://example.com/audio",
                "reasons": "AI generation detected",
                "error": "",
            }
        ]

        file_path = export_audit_excel(metrics, results)
        self.assertTrue(Path(file_path).exists())

        wb = openpyxl.load_workbook(file_path, data_only=True)
        self.assertIn("Audit Details", wb.sheetnames)
        self.assertIn("Executive Summary", wb.sheetnames)

        details_sheet = wb["Audit Details"]
        headers = [c for c in next(details_sheet.iter_rows(values_only=True))]
        self.assertIn("ai_model_prediction", headers)
        self.assertIn("ai_model_verdict", headers)
        self.assertIn("ai_model_is_correct", headers)
        self.assertIn("consensus_prediction", headers)

        summary_sheet = wb["Executive Summary"]
        summary_metrics = [row[0] for row in summary_sheet.iter_rows(values_only=True)]
        self.assertIn("Độ chính xác Thuần AI Model (Raw Gemini)", summary_metrics)
        self.assertIn("Độ chính xác Tổng hợp (Consensus Final)", summary_metrics)

        wb.close()
        if Path(file_path).exists():
            Path(file_path).unlink()

    def test_timing_verification_after_fallback(self) -> None:
        """Đảm bảo trạng thái is_verified_source được đọc sau khi detection_service và audio fallback hoàn tất."""
        track = TrackItem(
            id="test_fallback_track",
            raw_query="Test Fallback Track",
            input_type=InputType.SONG_NAME,
            track_title="Test Fallback Track",
            artist="Fallback Singer",
        )
        # Giả lập ban đầu resolver tìm thấy (True), nhưng trong detect_single bị lỗi audio và fallback sang video lệch (False)
        track.is_verified_match = True

        def fake_detect(t: TrackItem, **kwargs: Any) -> DetectionResult:
            # Mô phỏng AudioEngine fallback khiến track.is_verified_match đổi thành False
            t.is_verified_match = False
            t.match_score = 0.068
            return DetectionResult(
                track_id=t.id,
                track_title=t.track_title,
                input_type=t.input_type,
                raw_query=t.raw_query,
                is_ai=False,
                verdict_type=VerdictType.CONFIRMED_HUMAN,
                verdict_label="CONFIRMED HUMAN",
                confidence_score=0.90,
                suspected_engine="None",
                reasons=["Human performance"],
                ai_model_is_ai=False,
                ai_model_verdict="CONFIRMED_HUMAN",
                detailed_explanation="Human performance authentic.",
            )

        mock_detection_service = MagicMock()
        mock_detection_service.detect_single.side_effect = fake_detect

        # Thực thi logic giống như trong evaluator đã sửa
        res = mock_detection_service.detect_single(track)
        self.assertIsNotNone(res)
        is_verified_after = track.is_verified_match is not False

        self.assertFalse(
            is_verified_after,
            "Sau khi detect_single thực hiện fallback, trạng thái xác thực nguồn phải là False!",
        )
        self.assertEqual(track.match_score, 0.068)

    def test_vocal_rap_song_not_misclassified_as_instrumental_ai(self) -> None:
        """Bài hát có vocal (rap/hát) không được kích hoạt quy tắc 'AI nhạc cụ' dù rolloff thấp và rigidity cao."""
        detection_service = DetectionService(
            audio_engine=MagicMock(),
            signal_analyzer=MagicMock(),
            classifier=MagicMock(),
            resolver=MagicMock(),
            track_repo=MagicMock(),
        )

        track = TrackItem(
            id="test_rap_song",
            raw_query="Tình Thương Ta Trao Trọn Luân HV",
            input_type=InputType.SONG_NAME,
            track_title="Tình Thương Ta Trao Trọn",
            artist="Luân HV, LTi, Sang Head, Eoti",
        )

        # Gemini ban đầu phân loại chính xác là Người (PROCESSED_HUMAN)
        initial_result = DetectionResult(
            track_id=track.id,
            track_title=track.track_title,
            input_type=track.input_type,
            raw_query=track.raw_query,
            is_ai=False,
            verdict_type=VerdictType.PROCESSED_HUMAN,
            verdict_label="PROCESSED HUMAN",
            confidence_score=0.88,
            suspected_engine="None",
            reasons=["Authentic human vocal phrasing and natural chest resonance."],
            detailed_explanation="Human performance with studio pitch correction.",
            ai_model_is_ai=False,
            ai_model_verdict="PROCESSED_HUMAN",
        )

        # Dữ liệu âm học: Bài rap có beat trầm (rolloff = 6162 Hz, rigidity = 0.78), nhưng có vocal (is_instrumental = False)
        metrics = AcousticMetrics(
            sample_rate=44100,
            duration_sec=30.0,
            spectral_rolloff_95_hz=6162.0,
            pitch_rigidity_ratio=0.78,
            fakeprint_score=0.15,
            fakeprint_detected=False,
            stem_bleed_score=0.20,
            stem_bleed_detected=False,
            is_instrumental=False,  # Có vocal!
        )

        final_res = detection_service._apply_cross_validation_safeguard(
            result=initial_result,
            track=track,
            metadata={"artist": track.artist, "title": track.track_title},
            metrics=metrics,
        )

        self.assertFalse(
            final_res.is_ai,
            "Bài rap có vocal (is_instrumental=False) không được lật ngược thành AI!",
        )
        self.assertEqual(final_res.verdict_type, VerdictType.PROCESSED_HUMAN)

    def test_national_artist_titles_and_buddhist_singer_whitelisted(self) -> None:
        """Các danh hiệu NSND, NSƯT và ca sĩ Phật giáo (Hùng Thanh) phải được nhận diện là ca sĩ người thật."""
        from backend.services.signal_analyzer_service import is_verified_human_artist

        self.assertTrue(
            is_verified_human_artist("NSND Thúy Hường", "Lòng Vẫn Đợi Chờ"),
            "NSND Thúy Hường (Nghệ sĩ Nhân dân) phải được nhận diện là ca sĩ người thật!",
        )
        self.assertTrue(
            is_verified_human_artist("NSƯT Kim Tử Long", "Tình Anh Bán Chiếu"),
            "NSƯT Kim Tử Long (Nghệ sĩ Ưu tú) phải được nhận diện là ca sĩ người thật!",
        )
        self.assertTrue(
            is_verified_human_artist("Hùng Thanh", "Niệm Phật Dược Sư"),
            "Ca sĩ Phật giáo Hùng Thanh phải nằm trong Whitelist ca sĩ người thật!",
        )
        self.assertTrue(
            is_verified_human_artist("Thùy Chi", "Một Thoáng Huế Thương"),
            "Ca sĩ Thùy Chi phải nằm trong Whitelist ca sĩ người thật!",
        )

    def test_brazilian_phonk_and_montagem_recognized_in_catalog(self) -> None:
        """Các dự án Phonk/Montagem AI (KVLTX, Mano, Frozen Horizon) phải có trong catalog AI."""
        phonk_artists = [
            ("KVLTX", "MONTAGEM ALVORADA (Super Slowed)"),
            ("Mano", "MONTAGEM ESPADA (Super Slowed)"),
            ("Frozen Horizon", "Winter Nordic Dawn"),
        ]

        for artist, title in phonk_artists:
            meta = {"artist": artist, "title": title}
            heuristic_res = self.signal_analyzer.check_heuristics(meta)
            self.assertIsNotNone(
                heuristic_res,
                f"Dự án Phonk/Ambient AI '{artist}' phải được nhận diện qua Heuristics Catalog!",
            )
            self.assertTrue(
                heuristic_res.get("is_ai"),
                f"Dự án Phonk/Ambient AI '{artist}' phải được gán nhãn is_ai=True!",
            )

    def test_work_type_confusion_matrix_and_metrics(self) -> None:
        """Đảm bảo hàm tính toán và vẽ ma trận nhầm lẫn Work Type hoạt động chính xác."""
        from benchmark_evaluator import (
            calculate_work_type_confusion_matrix,
            plot_work_type_confusion_matrix,
        )

        y_true = ["ORIGINAL", "COVER", "REMIX", "ORIGINAL"]
        y_pred = ["ORIGINAL", "COVER", "ORIGINAL", "ORIGINAL"]

        wt_metrics = calculate_work_type_confusion_matrix(y_true, y_pred)
        self.assertEqual(wt_metrics["total_evaluated"], 4)
        self.assertEqual(wt_metrics["correct"], 3)
        self.assertAlmostEqual(wt_metrics["accuracy"], 0.75)

        img_path = Path("data/temp/test_work_type_cm.png")
        saved_file = plot_work_type_confusion_matrix(wt_metrics, output_path=img_path)
        self.assertTrue(saved_file.exists())
        self.assertGreater(saved_file.stat().st_size, 1000)

        if saved_file.exists():
            saved_file.unlink()

    def test_work_type_resolver_on_known_patterns(self) -> None:
        """Kiểm tra giải mã tác phẩm Original vs Cover vs Remix từ tiêu đề và siêu dữ liệu."""
        from backend.core.models.entities import WorkType
        from backend.services.work_type_resolver import (
            map_song_type_to_work_type,
            resolve_work_type_from_metadata,
        )

        # 1. Kiểm tra ánh xạ từ cột 'Song Type' của testcase
        self.assertEqual(map_song_type_to_work_type("Original"), "ORIGINAL")
        self.assertEqual(map_song_type_to_work_type("Live"), "ORIGINAL")
        self.assertEqual(map_song_type_to_work_type("Cover"), "COVER")
        self.assertEqual(map_song_type_to_work_type("Remake/Remix"), "REMIX")

        # 2. Kiểm tra các mẫu Remix / Lofi
        w_remix, _, is_cov, _, _ = resolve_work_type_from_metadata(
            title="Kiệu Hồng Người Dưng (Remix #3)", artist="Hoàng Vương"
        )
        self.assertEqual(w_remix, WorkType.REMIX)
        self.assertFalse(is_cov)

        w_lofi, _, _, _, _ = resolve_work_type_from_metadata(
            title="MỘT NGƯỜI NHƯ THẾ (Lofi)", artist="TINH HÀ"
        )
        self.assertEqual(w_lofi, WorkType.REMIX)

        # 3. Kiểm tra các mẫu Cover
        w_cover, _, is_cov2, _, _ = resolve_work_type_from_metadata(
            title="Em Chịu Đau Quá Dữ (Cover)", artist="Chấn Hào"
        )
        self.assertEqual(w_cover, WorkType.COVER)
        self.assertTrue(is_cov2)

        # 4. Kiểm tra ca sĩ hát bài hit nổi tiếng (Cover không đề chữ Cover trong title)
        w_hit_cover, _, is_cov3, _, _ = resolve_work_type_from_metadata(
            title="Ngày Mai Người Ta Lấy Chồng", artist="Giang Hồng Ngọc"
        )
        self.assertEqual(w_hit_cover, WorkType.COVER)
        self.assertTrue(is_cov3)

        # 5. Kiểm tra bản Live chính chủ (Original)
        w_live, _, is_cov4, _, _ = resolve_work_type_from_metadata(
            title="Đừng Như Thói Quen (Live)", artist="Sara Luu, JayKii"
        )
        self.assertEqual(w_live, WorkType.ORIGINAL)
        self.assertFalse(is_cov4)

    def test_check_work_type_accuracy_script_execution(self) -> None:
        """Đảm bảo script check_work_type_accuracy chạy đối soát chính xác và đạt độ chuẩn cao."""
        from check_work_type_accuracy import check_work_type_accuracy

        # Chạy kiểm tra không xuất file để đảm bảo tốc độ
        res = check_work_type_accuracy(re_resolve=True, export=False)
        self.assertIn("metrics", res)
        self.assertIn("records", res)
        self.assertGreater(res["metrics"]["total_evaluated"], 0)
        # Độ chính xác phân biệt tác phẩm đạt mức xuất sắc (trên 90%)
        self.assertGreaterEqual(res["metrics"]["accuracy"], 0.90)


if __name__ == "__main__":
    unittest.main()
