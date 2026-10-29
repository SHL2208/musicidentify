"""Bộ công cụ đánh giá độ chính xác & vẽ Ma trận nhầm lẫn (Confusion Matrix).

Hỗ trợ thẩm định mô hình AI Music Detector bằng dữ liệu Ground Truth, tính toán
đầy đủ các chỉ số chuẩn Machine Learning (Accuracy, Precision, Recall, Specificity,
F1-Score), trực quan hóa Confusion Matrix chất lượng cao (300 DPI) và tích hợp
bộ kiểm thử tự động (Unit Test).

Tuân thủ nghiêm ngặt chuẩn PEP 8 và Ruff.
"""

from __future__ import annotations

import argparse
import os
import random
import sys
import time
import unittest
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, cast

if sys.platform == "win32":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")


# Thêm thư mục gốc dự án vào sys.path
ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

# Matplotlib backend headless để hoạt động ổn định trên cả môi trường CI/Terminal
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd

from benchmark_evaluator import (
    map_song_type_to_work_type,
)


@dataclass
class ConfusionMetrics:
    """Chứa toàn bộ các chỉ số thống kê của ma trận nhầm lẫn."""

    total_samples: int
    evaluated_samples: int
    errors: int
    tp: int  # True Positive (Thực tế AI, Đoán AI)
    tn: int  # True Negative (Thực tế Human, Đoán Human)
    fp: int  # False Positive (Thực tế Human, Đoán AI - Báo động giả)
    fn: int  # False Negative (Thực tế AI, Đoán Human - Lọt lưới)
    accuracy: float
    precision: float
    recall: float
    specificity: float
    f1_score: float
    balanced_accuracy: float
    error_rate: float
    error_records: list[dict[str, Any]] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        """Chuyển đổi dữ liệu thống kê thành dictionary."""
        return asdict(self)


def calculate_confusion_metrics(
    y_true: list[bool],
    y_pred: list[bool],
    error_count: int = 0,
    error_records: list[dict[str, Any]] | None = None,
) -> ConfusionMetrics:
    """Tính toán ma trận nhầm lẫn và các chỉ số đo lường hiệu năng.

    Parameters:
        y_true: Danh sách nhãn thực tế (True = AI, False = Human).
        y_pred: Danh sách nhãn dự đoán (True = AI, False = Human).
        error_count: Số lượng mẫu gặp sự cố kết nối/xử lý lỗi.
        error_records: Danh sách chi tiết các bài hát gặp lỗi và bài nạp bù.

    Returns:
        ConfusionMetrics: Đối tượng chứa đầy đủ các chỉ số thống kê.
    """
    if len(y_true) != len(y_pred):
        raise ValueError(f"Độ dài y_true ({len(y_true)}) và y_pred ({len(y_pred)}) phải bằng nhau.")

    tp = sum(1 for yt, yp in zip(y_true, y_pred, strict=True) if yt and yp)
    tn = sum(1 for yt, yp in zip(y_true, y_pred, strict=True) if not yt and not yp)
    fp = sum(1 for yt, yp in zip(y_true, y_pred, strict=True) if not yt and yp)
    fn = sum(1 for yt, yp in zip(y_true, y_pred, strict=True) if yt and not yp)

    evaluated = tp + tn + fp + fn
    total = evaluated + error_count

    accuracy = (tp + tn) / evaluated if evaluated > 0 else 0.0
    precision = tp / (tp + fp) if (tp + fp) > 0 else 0.0
    recall = tp / (tp + fn) if (tp + fn) > 0 else 0.0
    specificity = tn / (tn + fp) if (tn + fp) > 0 else 0.0
    f1_score = (2 * precision * recall) / (precision + recall) if (precision + recall) > 0 else 0.0
    balanced_acc = (recall + specificity) / 2 if evaluated > 0 else 0.0
    error_rate = (fp + fn) / evaluated if evaluated > 0 else 0.0

    return ConfusionMetrics(
        total_samples=total,
        evaluated_samples=evaluated,
        errors=error_count,
        tp=tp,
        tn=tn,
        fp=fp,
        fn=fn,
        accuracy=accuracy,
        precision=precision,
        recall=recall,
        specificity=specificity,
        f1_score=f1_score,
        balanced_accuracy=balanced_acc,
        error_rate=error_rate,
        error_records=error_records or [],
    )


def plot_confusion_matrix(
    metrics: ConfusionMetrics,
    output_path: str | Path | None = None,
    title: str = "Confusion Matrix & Accuracy Assessment",
    model_name: str | None = None,
) -> Path:
    """Vẽ ma trận nhầm lẫn 2x2 chuyên nghiệp và bảng thống kê chỉ số hiệu năng.

    Parameters:
        metrics: Đối tượng ConfusionMetrics chứa các chỉ số.
        output_path: Đường dẫn lưu file ảnh (mặc định trong data/outputs).
        title: Tiêu đề biểu đồ.
        model_name: Tên mô hình hoặc ghi chú đánh giá.

    Returns:
        Path: Đường dẫn file ảnh PNG đã lưu.
    """
    if output_path is None:
        out_dir = ROOT_DIR / "data" / "outputs"
        out_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        target_file = out_dir / f"confusion_matrix_{ts}.png"
    else:
        target_file = Path(output_path)
        target_file.parent.mkdir(parents=True, exist_ok=True)

    actual_human_total = metrics.tn + metrics.fp
    actual_ai_total = metrics.fn + metrics.tp

    tn_pct = (metrics.tn / actual_human_total * 100) if actual_human_total > 0 else 0.0
    fp_pct = (metrics.fp / actual_human_total * 100) if actual_human_total > 0 else 0.0
    fn_pct = (metrics.fn / actual_ai_total * 100) if actual_ai_total > 0 else 0.0
    tp_pct = (metrics.tp / actual_ai_total * 100) if actual_ai_total > 0 else 0.0

    # Khởi tạo khung hình 2 panel (Confusion Matrix bên trái, Bảng chỉ số bên phải)
    fig, (ax_cm, ax_stats) = plt.subplots(
        nrows=1,
        ncols=2,
        figsize=(14, 6.8),
        dpi=300,
        gridspec_kw={"width_ratios": [1.15, 1.0]},
    )
    fig.patch.set_facecolor("#FAFAFC")

    # 1. Vẽ Heatmap bên trái
    cmap = plt.cm.Blues
    norm_matrix = np.array(
        [
            [tn_pct / 100.0, fp_pct / 100.0],
            [fn_pct / 100.0, tp_pct / 100.0],
        ]
    )

    im = ax_cm.imshow(norm_matrix, interpolation="nearest", cmap=cmap, vmin=0.0, vmax=1.0)
    plt.colorbar(im, ax=ax_cm, fraction=0.046, pad=0.04, label="Tỉ lệ theo nhóm thực tế (%)")

    # Nhãn trục
    classes = ["HUMAN", "AI"]
    tick_marks = np.arange(len(classes))
    ax_cm.set_xticks(tick_marks)
    ax_cm.set_xticklabels([f"Đoán: {c}" for c in classes], fontsize=11, fontweight="bold")
    ax_cm.set_yticks(tick_marks)
    ax_cm.set_yticklabels([f"Thực tế: {c}" for c in classes], fontsize=11, fontweight="bold")
    ax_cm.set_xlabel("DỰ ĐOÁN CỦA HỆ THỐNG (PREDICTED)", fontsize=11, labelpad=10)
    ax_cm.set_ylabel("NHÃN THỰC TẾ (GROUND TRUTH)", fontsize=11, labelpad=10)
    ax_cm.set_title("MA TRẬN NHẦM LẪN (2x2 MATRIX)", fontsize=13, fontweight="bold", pad=12)

    # Điền số liệu chi tiết từng ô
    cell_info = [
        [
            {
                "label": "TRUE NEGATIVE (TN)",
                "count": metrics.tn,
                "pct": tn_pct,
                "desc": "✓ Đúng: Nhạc người",
                "is_correct": True,
            },
            {
                "label": "FALSE POSITIVE (FP)",
                "count": metrics.fp,
                "pct": fp_pct,
                "desc": "✗ Báo động giả (Vu oan)",
                "is_correct": False,
            },
        ],
        [
            {
                "label": "FALSE NEGATIVE (FN)",
                "count": metrics.fn,
                "pct": fn_pct,
                "desc": "✗ Lọt lưới (Bỏ sót AI)",
                "is_correct": False,
            },
            {
                "label": "TRUE POSITIVE (TP)",
                "count": metrics.tp,
                "pct": tp_pct,
                "desc": "✓ Đúng: Nhạc AI",
                "is_correct": True,
            },
        ],
    ]

    for i in range(2):
        for j in range(2):
            info = cell_info[i][j]
            val_pct = norm_matrix[i, j]
            text_color = "white" if val_pct > 0.55 else "#1A202C"

            txt = f"{info['label']}\nSố lượng: {info['count']} ({info['pct']:.1f}%)\n{info['desc']}"
            ax_cm.text(
                j,
                i,
                txt,
                ha="center",
                va="center",
                color=text_color,
                fontsize=10.5,
                fontweight="bold",
                linespacing=1.4,
            )

    # 2. Vẽ bảng chỉ số và biểu đồ đo lường bên phải (ax_stats)
    ax_stats.set_facecolor("#FFFFFF")
    ax_stats.spines["top"].set_visible(False)
    ax_stats.spines["right"].set_visible(False)
    ax_stats.spines["left"].set_visible(False)
    ax_stats.spines["bottom"].set_color("#E2E8F0")

    metrics_list = [
        ("Accuracy (Độ chính xác)", metrics.accuracy),
        ("Precision (Độ tin cậy báo AI)", metrics.precision),
        ("Recall (Độ nhạy bắt AI)", metrics.recall),
        ("Specificity (Bảo vệ nhạc người)", metrics.specificity),
        ("F1-Score (Cân bằng P/R)", metrics.f1_score),
        ("Balanced Accuracy", metrics.balanced_accuracy),
    ]

    y_pos = np.arange(len(metrics_list))
    names = [m[0] for m in metrics_list]
    values = [m[1] * 100 for m in metrics_list]

    # Màu sắc thanh đo: Xanh nếu >= 80%, Cam nếu 60-80%, Đỏ nếu < 60%
    bar_colors = ["#10B981" if v >= 80.0 else "#F59E0B" if v >= 60.0 else "#EF4444" for v in values]

    bars = ax_stats.barh(
        y_pos,
        values,
        height=0.55,
        color=bar_colors,
        edgecolor="#CBD5E1",
        alpha=0.9,
    )
    ax_stats.set_yticks(y_pos)
    ax_stats.set_yticklabels(names, fontsize=10.5, fontweight="bold")
    ax_stats.set_xlim(0, 118)
    ax_stats.set_xlabel("Tỷ lệ phần trăm (%)", fontsize=11, labelpad=8)
    ax_stats.set_title("BẢNG CHỈ SỐ ĐO LƯỜNG HIỆU NĂNG", fontsize=13, fontweight="bold", pad=12)
    ax_stats.invert_yaxis()

    # Thêm giá trị cụ thể tại đầu mỗi thanh bar
    for bar, val in zip(bars, values, strict=True):
        width = bar.get_width()
        ax_stats.text(
            width + 2,
            bar.get_y() + bar.get_height() / 2,
            f"{val:.1f}%",
            va="center",
            ha="left",
            fontsize=10.5,
            fontweight="bold",
            color="#1E293B",
        )

    # Thêm hộp thông tin tổng kết bên dưới
    summary_text = (
        f"• Tổng số mẫu kiểm thử: {metrics.total_samples} bài "
        f"(Thành công: {metrics.evaluated_samples} | Lỗi: {metrics.errors})\n"
        f"• Dự đoán đúng: {metrics.tp + metrics.tn} bài ({metrics.accuracy * 100:.1f}%) "
        f"| Dự đoán sai: {metrics.fp + metrics.fn} bài ({metrics.error_rate * 100:.1f}%)\n"
        f"• Báo động giả (FP): {metrics.fp} bài | Bỏ sót AI (FN): {metrics.fn} bài\n"
        f"• Mô hình đánh giá: {model_name or 'Project Core AI Engine'}"
    )

    fig.text(
        0.5,
        0.025,
        summary_text,
        ha="center",
        fontsize=9.8,
        bbox={
            "boxstyle": "round,pad=0.6",
            "facecolor": "#F1F5F9",
            "edgecolor": "#CBD5E1",
            "alpha": 0.9,
        },
    )

    # Tiêu đề tổng quát phía trên
    fig.suptitle(title, fontsize=15, fontweight="bold", y=0.98, color="#0F172A")
    plt.tight_layout(rect=(0.02, 0.12, 0.98, 0.93))

    plt.savefig(target_file, dpi=300, facecolor=fig.get_facecolor(), bbox_inches="tight")
    plt.close(fig)

    return target_file


def load_testcase_dataset(
    file_path: str | Path = "testcase.xlsx",
    sheet_name: str = "Media_W38",
) -> list[dict[str, Any]]:
    """Đọc dữ liệu bài hát kèm nhãn Ground Truth từ file Excel testcase."""
    path_obj = Path(file_path)
    if not path_obj.is_absolute():
        if (ROOT_DIR / path_obj).exists():
            path_obj = ROOT_DIR / path_obj
        elif (ROOT_DIR / "data" / path_obj).exists():
            path_obj = ROOT_DIR / "data" / path_obj
        else:
            path_obj = ROOT_DIR / path_obj

    if not path_obj.exists():
        raise FileNotFoundError(f"Không tìm thấy file dataset: {path_obj}")

    wb = openpyxl.load_workbook(str(path_obj), data_only=True)
    active_sheet = sheet_name if sheet_name in wb.sheetnames else wb.sheetnames[0]
    sheet = wb[active_sheet]

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Sheet '{active_sheet}' không có dữ liệu.")

    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]

    title_idx = -1
    artist_idx = -1
    ai_col_idx = -1
    media_id_idx = -1
    note_idx = -1
    song_type_idx = -1
    release_date_idx = -1
    distributor_idx = -1

    for idx, h in enumerate(headers):
        if any(k in h for k in ["song title", "tên bài", "title"]):
            title_idx = idx
        elif any(k in h for k in ["artist", "ca sĩ", "nghệ sĩ"]):
            artist_idx = idx
        elif any(k in h for k in ["ai yes", "ai"]):
            ai_col_idx = idx
        elif any(k in h for k in ["media_id", "id"]):
            media_id_idx = idx
        elif any(k in h for k in ["note", "ghi chú"]):
            note_idx = idx
        elif any(k in h for k in ["song type", "loại bài", "type"]):
            song_type_idx = idx
        elif any(k in h for k in ["release", "ngày phát hành"]):
            release_date_idx = idx
        elif any(k in h for k in ["distributor", "nhà phân phối"]):
            distributor_idx = idx

    if title_idx == -1:
        title_idx = 1 if len(headers) > 1 else 0

    dataset: list[dict[str, Any]] = []
    for r_idx, r in enumerate(rows[1:], start=2):
        title = str(r[title_idx]).strip() if r[title_idx] is not None else ""
        if not title:
            continue

        artist = (
            str(r[artist_idx]).strip() if (artist_idx != -1 and r[artist_idx] is not None) else ""
        )
        media_id = (
            str(r[media_id_idx]).strip()
            if (media_id_idx != -1 and r[media_id_idx] is not None)
            else f"row_{r_idx}"
        )
        note = str(r[note_idx]).strip() if (note_idx != -1 and r[note_idx] is not None) else ""
        song_type = (
            str(r[song_type_idx]).strip()
            if (song_type_idx != -1 and r[song_type_idx] is not None)
            else ""
        )
        release_date = (
            str(r[release_date_idx]).strip()
            if (release_date_idx != -1 and r[release_date_idx] is not None)
            else ""
        )
        distributor = (
            str(r[distributor_idx]).strip()
            if (distributor_idx != -1 and r[distributor_idx] is not None)
            else ""
        )

        ai_val = (
            str(r[ai_col_idx]).strip().lower()
            if (ai_col_idx != -1 and r[ai_col_idx] is not None)
            else ""
        )
        is_ai_ground_truth = "yes" in ai_val or ai_val == "y"

        query = f"{title} {artist}".strip()
        dataset.append(
            {
                "media_id": media_id,
                "title": title,
                "artist": artist,
                "query": query,
                "gt_is_ai": is_ai_ground_truth,
                "gt_label": "AI" if is_ai_ground_truth else "HUMAN",
                "note": note,
                "song_type": song_type,
                "release_date": release_date,
                "distributor": distributor,
                "row_idx": r_idx,
            }
        )

    return dataset


def select_samples(
    dataset: list[dict[str, Any]],
    sample_size: int | None = 20,
    balanced: bool = True,
    seed: int = 42,
) -> list[dict[str, Any]]:
    """Lấy mẫu đánh giá từ dataset (hỗ trợ phân tầng cân bằng 50% AI - 50% Con người)."""
    if sample_size is None or sample_size >= len(dataset):
        return dataset

    random.seed(seed)
    ai_items = [d for d in dataset if d["gt_is_ai"]]
    human_items = [d for d in dataset if not d["gt_is_ai"]]

    if balanced:
        half = sample_size // 2
        selected_ai = random.sample(ai_items, min(half, len(ai_items)))
        remaining = sample_size - len(selected_ai)
        selected_human = random.sample(human_items, min(remaining, len(human_items)))
        combined = selected_ai + selected_human
        random.shuffle(combined)
        return combined

    return random.sample(dataset, sample_size)


def evaluate_dataset_with_pipeline(
    samples: list[dict[str, Any]],
    dataset: list[dict[str, Any]] | None = None,
    model_override: str | None = None,
    mock_mode: bool = False,
    mock_accuracy: float = 0.70,
) -> tuple[ConfusionMetrics, list[dict[str, Any]]]:
    """Chạy thẩm định toàn bộ tập mẫu qua hệ thống và trả về ma trận kết quả.

    Hỗ trợ cơ chế Nạp Bù Mẫu Tự Động (Resilient Sampling Pool):
    Khi 1 bài hát gặp lỗi tải/Zing VIP, tự động lấy bài kế tiếp cùng nhãn từ dataset dự phòng
    để đảm bảo đủ chỉ tiêu số bài thành công, đồng thời lưu vết toàn bộ bài lỗi.
    Hỗ trợ chế độ mock_mode để kiểm thử thuật toán nhanh chóng mà không tiêu tốn API.
    """
    results: list[dict[str, Any]] = []
    error_records: list[dict[str, Any]] = []
    y_true: list[bool] = []
    y_pred: list[bool] = []
    errors = 0

    if mock_mode:
        random.seed(42)
        for item in samples:
            gt_is_ai = item["gt_is_ai"]
            y_true.append(gt_is_ai)

            # Mô phỏng độ chính xác mock_accuracy (ví dụ 70%)
            is_correct = random.random() < mock_accuracy
            pred_is_ai = gt_is_ai if is_correct else not gt_is_ai
            y_pred.append(pred_is_ai)

            gt_work = map_song_type_to_work_type(item.get("song_type"))
            pred_work = gt_work if is_correct else "UNCERTAIN"
            work_is_correct = (gt_work == pred_work) if gt_work != "UNCERTAIN" else None

            results.append(
                {
                    "media_id": item["media_id"],
                    "song_title": item["title"],
                    "artist": item["artist"],
                    "query": item["query"],
                    "ground_truth": item["gt_label"],
                    "gt_work_type": gt_work,
                    "pred_work_type": pred_work,
                    "work_is_correct": "ĐÚNG"
                    if work_is_correct is True
                    else ("SAI" if work_is_correct is False else "N/A"),
                    "work_label": "Bản gốc (Original)" if pred_work == "ORIGINAL" else pred_work,
                    "original_artist": item["artist"],
                    "original_song_title": item["title"],
                    "is_cover": pred_work == "COVER",
                    "ai_model_prediction": "AI" if pred_is_ai else "HUMAN",
                    "ai_model_verdict": "AI_GENERATED" if pred_is_ai else "CONFIRMED_HUMAN",
                    "ai_model_is_correct": "ĐÚNG" if is_correct else "SAI",
                    "consensus_prediction": "AI" if pred_is_ai else "HUMAN",
                    "consensus_verdict": "AI_GENERATED" if pred_is_ai else "CONFIRMED_HUMAN",
                    "prediction": "AI" if pred_is_ai else "HUMAN",
                    "is_correct": "ĐÚNG" if is_correct else "SAI",
                    "confidence_score": "75.0%",
                    "reasons": "Mock detection evaluation",
                }
            )
    else:
        from backend.core.container import get_container

        container = get_container()
        detection_service = container.detection_service
        resolver = container.resolver

        total_target = len(samples)
        evaluation_queue = list(samples)
        attempted_ids = {s["media_id"] for s in samples}

        print("=" * 76)
        print(f"🚀 BẮT ĐẦU CHẤM ĐIỂM TESTCASE (MỤC TIÊU: {total_target} BÀI HÁT THÀNH CÔNG)")
        print("=" * 76)

        i = 0
        while i < len(evaluation_queue):
            item = evaluation_queue[i]
            i += 1
            gt_is_ai = item["gt_is_ai"]
            query = item["query"]
            gt_lbl = item["gt_label"]

            print(f"\n[{i}/{len(evaluation_queue)}] Thẩm định: '{query}' | Ground Truth: {gt_lbl}")
            try:
                track = resolver.resolve_track(item["title"], artist=item.get("artist", ""))
                track.song_type = item.get("song_type") or ""
                track.release_date = str(item.get("release_date") or "")
                track.distributor = str(item.get("distributor") or "")
                res = detection_service.detect_single(track, model_override=model_override)

                # Cập nhật trạng thái xác thực nguồn sau khi audio_engine đã chạy và xử lý fallback
                is_verified = track.is_verified_match is not False
                if not is_verified:
                    print(
                        f"       -> ⚠️ CẢNH BÁO NGUỒN: Audio không khớp chính xác nghệ sĩ! "
                        f"(Score: {track.match_score})"
                    )

                pred_is_ai = res.is_ai
                y_true.append(gt_is_ai)
                y_pred.append(pred_is_ai)
                is_correct = pred_is_ai == gt_is_ai

                if gt_is_ai and pred_is_ai:
                    cat = "TP (Đúng: Phát hiện AI)"
                elif not gt_is_ai and not pred_is_ai:
                    cat = "TN (Đúng: Xác nhận Người)"
                elif not gt_is_ai and pred_is_ai:
                    cat = "FP (Sai: Vu oan Người)"
                else:
                    cat = "FN (Sai: Bỏ sót AI)"

                icon = "✅" if is_correct else "❌"
                pred_label = "AI" if pred_is_ai else "HUMAN"
                match_status_text = "ĐÃ XÁC THỰC" if is_verified else "SAI LỆCH (UNVERIFIED)"

                ai_model_is_ai = getattr(res, "ai_model_is_ai", res.is_ai)
                ai_model_pred_label = "AI" if ai_model_is_ai else "HUMAN"
                ai_model_verdict_str = getattr(res, "ai_model_verdict", None) or (
                    res.verdict_type.value
                    if hasattr(res.verdict_type, "value")
                    else str(res.verdict_type)
                )
                ai_model_is_correct = "ĐÚNG" if ai_model_is_ai == gt_is_ai else "SAI"

                consensus_pred_label = pred_label
                consensus_verdict_str = getattr(res, "consensus_verdict", None) or (
                    res.verdict_type.value
                    if hasattr(res.verdict_type, "value")
                    else str(res.verdict_type)
                )

                gt_work = map_song_type_to_work_type(item.get("song_type"))
                pred_work = (
                    res.work_type.value
                    if hasattr(res.work_type, "value")
                    else str(getattr(res, "work_type", "UNCERTAIN"))
                )
                work_is_correct = (gt_work == pred_work) if gt_work != "UNCERTAIN" else None

                print(
                    f"       -> Model AI: {ai_model_verdict_str} | Tổng hợp: {consensus_verdict_str} ({res.verdict_label}) | {icon} {cat}"
                )
                print(
                    f"       -> Nguồn: {match_status_text} | "
                    f"Độ tin cậy: {res.confidence_score * 100:.1f}% | "
                    f"Lý do: {res.reasons[0] if res.reasons else 'N/A'}"
                )

                results.append(
                    {
                        "media_id": item["media_id"],
                        "song_title": item["title"],
                        "artist": item["artist"],
                        "query": query,
                        "ground_truth": gt_lbl,
                        "gt_work_type": gt_work,
                        "pred_work_type": pred_work,
                        "work_is_correct": "ĐÚNG"
                        if work_is_correct is True
                        else ("SAI" if work_is_correct is False else "N/A"),
                        "work_label": getattr(res, "work_label", "Chưa rõ tác phẩm"),
                        "original_artist": getattr(res, "original_artist", "") or "N/A",
                        "original_song_title": getattr(res, "original_song_title", "") or "N/A",
                        "is_cover": getattr(res, "is_cover", False),
                        "ai_model_prediction": ai_model_pred_label,
                        "ai_model_verdict": ai_model_verdict_str,
                        "ai_model_is_correct": ai_model_is_correct,
                        "consensus_prediction": consensus_pred_label,
                        "consensus_verdict": consensus_verdict_str,
                        "prediction": pred_label,
                        "is_correct": "ĐÚNG" if is_correct else "SAI",
                        "classification_type": cat,
                        "is_verified_source": is_verified,
                        "source_verification": match_status_text,
                        "match_score": (
                            f"{track.match_score * 100:.1f}%"
                            if track.match_score is not None
                            else "N/A"
                        ),
                        "audio_url": track.resolved_url or "",
                        "note": item.get("note", ""),
                        "verdict_type": res.verdict_type.value
                        if hasattr(res.verdict_type, "value")
                        else str(res.verdict_type),
                        "verdict_label": res.verdict_label,
                        "confidence_score": f"{res.confidence_score * 100:.1f}%",
                        "suspected_engine": res.suspected_engine,
                        "reasons": " | ".join(res.reasons or []),
                    }
                )
            except Exception as e:  # noqa: BLE001
                errors += 1
                print(f"       -> ❌ LỖI XỬ LÝ: {e}")

                # Tìm mẫu thay thế cùng nhãn từ dataset dự phòng
                replacement: dict[str, Any] | None = None
                if dataset:
                    for cand in dataset:
                        if (
                            cand["gt_is_ai"] == item["gt_is_ai"]
                            and cand["media_id"] not in attempted_ids
                        ):
                            replacement = cand
                            attempted_ids.add(cand["media_id"])
                            evaluation_queue.append(cand)
                            break

                error_records.append(
                    {
                        "Mã bài lỗi (Media ID)": item["media_id"],
                        "Tên bài lỗi (Song Title)": item["title"],
                        "Ca sĩ (Artist)": item["artist"],
                        "Nhãn thực tế (Ground Truth)": gt_lbl,
                        "Nguyên nhân sự cố (Error)": str(e),
                        "Bài nạp bù thay thế (Replacement)": replacement["query"]
                        if replacement
                        else "Không còn bài dự phòng",
                        "Thời điểm ghi nhận": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    }
                )

                if replacement:
                    print(
                        f"       -> 🔄 ĐÃ NẠP BÙ MẪU THAY THẾ: '{replacement['query']}' (cùng nhãn {gt_lbl})"
                    )

                results.append(
                    {
                        "media_id": item["media_id"],
                        "song_title": item["title"],
                        "artist": item["artist"],
                        "query": query,
                        "ground_truth": gt_lbl,
                        "gt_work_type": map_song_type_to_work_type(item.get("song_type")),
                        "pred_work_type": "ERROR",
                        "work_is_correct": "LỖI",
                        "work_label": "LỖI HỆ THỐNG",
                        "original_artist": "N/A",
                        "original_song_title": "N/A",
                        "is_cover": False,
                        "ai_model_prediction": "ERROR",
                        "ai_model_verdict": "ERROR",
                        "ai_model_is_correct": "LỖI",
                        "consensus_prediction": "ERROR",
                        "consensus_verdict": "ERROR",
                        "prediction": "ERROR",
                        "is_correct": "LỖI",
                        "classification_type": "ERROR",
                        "verdict_type": "ERROR",
                        "verdict_label": "LỖI HỆ THỐNG",
                        "confidence_score": "0%",
                        "reasons": str(e),
                    }
                )
            finally:
                time.sleep(2.5)

    metrics = calculate_confusion_metrics(
        y_true, y_pred, error_count=errors, error_records=error_records
    )
    return metrics, results


from backend.core.utils import clean_excel_text


def export_evaluation_excel(
    metrics: ConfusionMetrics,
    results: list[dict[str, Any]],
    output_path: str | Path | None = None,
) -> Path:
    """Xuất file Excel đối soát chi tiết kết quả thẩm định và ma trận nhầm lẫn."""
    if output_path is None:
        out_dir = ROOT_DIR / "data" / "outputs"
        out_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        target_file = out_dir / f"evaluation_audit_{ts}.xlsx"
    else:
        target_file = Path(output_path)
        target_file.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(str(target_file), engine="openpyxl") as writer:
        df_details = pd.DataFrame(results)
        df_any = cast(Any, df_details)
        if hasattr(df_any, "map"):
            df_details = df_any.map(clean_excel_text)
        elif hasattr(df_any, "applymap"):
            df_details = df_any.applymap(clean_excel_text)
        df_details.to_excel(writer, sheet_name="Audit Details", index=False)

        ai_model_correct = sum(1 for r in results if r.get("ai_model_is_correct") == "ĐÚNG")
        ai_model_total = sum(1 for r in results if r.get("ai_model_is_correct") in ("ĐÚNG", "SAI"))
        ai_model_acc_str = (
            f"{(ai_model_correct / ai_model_total * 100):.2f}%" if ai_model_total > 0 else "N/A"
        )

        summary_rows = [
            {"Chỉ số": "Tổng số bài mẫu", "Giá trị": metrics.total_samples},
            {
                "Chỉ số": "Số bài thẩm định thành công",
                "Giá trị": metrics.evaluated_samples,
            },
            {"Chỉ số": "Số bài gặp sự cố", "Giá trị": metrics.errors},
            {"Chỉ số": "Số bài lỗi đã nạp bù", "Giá trị": len(metrics.error_records)},
            {
                "Chỉ số": "Độ chính xác Thuần AI Model (Raw Gemini)",
                "Giá trị": ai_model_acc_str,
            },
            {
                "Chỉ số": "Độ chính xác Tổng hợp (Consensus Final)",
                "Giá trị": f"{metrics.accuracy * 100:.2f}%",
            },
            {"Chỉ số": "True Positive (Đoán đúng AI)", "Giá trị": metrics.tp},
            {"Chỉ số": "True Negative (Đoán đúng Người)", "Giá trị": metrics.tn},
            {"Chỉ số": "False Positive (Vu oan Người)", "Giá trị": metrics.fp},
            {"Chỉ số": "False Negative (Bỏ sót AI)", "Giá trị": metrics.fn},
            {
                "Chỉ số": "Accuracy (Độ chính xác)",
                "Giá trị": f"{metrics.accuracy * 100:.2f}%",
            },
            {
                "Chỉ số": "Precision (Độ tin cậy AI)",
                "Giá trị": f"{metrics.precision * 100:.2f}%",
            },
            {
                "Chỉ số": "Recall (Độ nhạy bắt AI)",
                "Giá trị": f"{metrics.recall * 100:.2f}%",
            },
            {
                "Chỉ số": "Specificity (Đặc hiệu nhạc người)",
                "Giá trị": f"{metrics.specificity * 100:.2f}%",
            },
            {"Chỉ số": "F1-Score", "Giá trị": f"{metrics.f1_score * 100:.2f}%"},
            {
                "Chỉ số": "Balanced Accuracy",
                "Giá trị": f"{metrics.balanced_accuracy * 100:.2f}%",
            },
            {
                "Chỉ số": "Error Rate (Tỷ lệ sai)",
                "Giá trị": f"{metrics.error_rate * 100:.2f}%",
            },
        ]
        df_summary = pd.DataFrame(summary_rows)
        df_summary.to_excel(writer, sheet_name="Executive Summary", index=False)

        # Sheet 3: Sự cố & Nạp bù (Error Audit & Replacement)
        if metrics.error_records:
            df_errors = pd.DataFrame(metrics.error_records)
            df_err_any = cast(Any, df_errors)
            if hasattr(df_err_any, "map"):
                df_errors = df_err_any.map(clean_excel_text)
            elif hasattr(df_err_any, "applymap"):
                df_errors = df_err_any.applymap(clean_excel_text)
            df_errors.to_excel(writer, sheet_name="Sự cố & Nạp bù", index=False)

    return target_file


def load_metrics_from_audit_report(
    report_file_path: str | Path,
) -> tuple[ConfusionMetrics, list[dict[str, Any]]]:
    """Đọc dữ liệu đối soát từ file Excel kết quả kiểm tra trước đó để vẽ Confusion Matrix.

    Giúp người dùng tận dụng kết quả đã test (ví dụ 70%) để vẽ biểu đồ ngay lập tức
    mà không cần phải chạy lại toàn bộ mô hình tốn thời gian.
    """
    path_obj = Path(report_file_path)
    if not path_obj.is_absolute():
        path_obj = ROOT_DIR / path_obj

    if not path_obj.exists():
        raise FileNotFoundError(f"Không tìm thấy file báo cáo: {path_obj}")

    with pd.ExcelFile(str(path_obj)) as xl:
        target_sheet = "Audit Details" if "Audit Details" in xl.sheet_names else xl.sheet_names[0]
        df = pd.read_excel(xl, sheet_name=target_sheet)

    # Tìm các cột Ground Truth và Prediction cho phân loại AI
    gt_col = None
    pred_col = None

    # 1. Tìm cột Ground Truth (ưu tiên nhãn AI thực tế, tránh các cột tác phẩm như gt_work_type)
    gt_exact = [
        "ground_truth",
        "ground_truth_ai",
        "gt_label",
        "nhãn thực tế",
        "thực tế",
        "gt_is_ai",
        "is_ai",
    ]
    for cand in gt_exact:
        for col in df.columns:
            if str(col).strip().lower() == cand:
                gt_col = col
                break
        if gt_col:
            break

    if gt_col is None:
        for col in df.columns:
            c_low = str(col).strip().lower()
            if (
                ("ground" in c_low or "thực tế" in c_low or c_low == "gt")
                and "work" not in c_low
                and "type" not in c_low
            ):
                gt_col = col
                break

    # 2. Tìm cột Prediction (ưu tiên kết quả dự đoán AI/Human, tránh pred_work_type hay các cột is_correct)
    pred_exact = [
        "prediction",
        "consensus_prediction",
        "ai_model_prediction",
        "dự đoán",
        "kết quả dự đoán",
    ]
    for cand in pred_exact:
        for col in df.columns:
            if str(col).strip().lower() == cand:
                pred_col = col
                break
        if pred_col:
            break

    if pred_col is None:
        for col in df.columns:
            c_low = str(col).strip().lower()
            if (
                ("prediction" in c_low or "dự đoán" in c_low or c_low == "pred")
                and "work" not in c_low
                and "type" not in c_low
                and "is_correct" not in c_low
            ):
                pred_col = col
                break

    if gt_col is None or pred_col is None:
        raise ValueError(
            f"Không tìm thấy cột Ground Truth hoặc Prediction trong {path_obj}. "
            f"Các cột hiện có: {list(df.columns)}"
        )

    y_true: list[bool] = []
    y_pred: list[bool] = []
    errors = 0
    records: list[dict[str, Any]] = []

    for _, row in df.iterrows():
        gt_val = str(row[gt_col]).strip().upper()
        pred_val = str(row[pred_col]).strip().upper()

        if "ERROR" in pred_val or "LỖI" in pred_val:
            errors += 1
            continue

        gt_is_ai = "AI" in gt_val or "YES" in gt_val
        pred_is_ai = "AI" in pred_val or "YES" in pred_val

        y_true.append(gt_is_ai)
        y_pred.append(pred_is_ai)
        records.append(dict(row))

    metrics = calculate_confusion_metrics(y_true, y_pred, error_count=errors)
    return metrics, records


# ==============================================================================
# BỘ UNIT TEST TỰ ĐỘNG (UNITTEST TEST SUITE)
# ==============================================================================


class TestConfusionMatrixEvaluation(unittest.TestCase):
    """Kiểm thử tự động tính toán chỉ số và vẽ biểu đồ ma trận nhầm lẫn."""

    def test_perfect_accuracy(self) -> None:
        """Kiểm tra trường hợp dự đoán chính xác 100%."""
        y_true = [True, True, False, False]
        y_pred = [True, True, False, False]
        metrics = calculate_confusion_metrics(y_true, y_pred)

        self.assertEqual(metrics.tp, 2)
        self.assertEqual(metrics.tn, 2)
        self.assertEqual(metrics.fp, 0)
        self.assertEqual(metrics.fn, 0)
        self.assertAlmostEqual(metrics.accuracy, 1.0)
        self.assertAlmostEqual(metrics.precision, 1.0)
        self.assertAlmostEqual(metrics.recall, 1.0)
        self.assertAlmostEqual(metrics.f1_score, 1.0)

    def test_seventy_percent_accuracy_scenario(self) -> None:
        """Kiểm tra trực tiếp kịch bản người dùng phản ánh: Tỉ lệ đạt ~70%.

        Ví dụ tập mẫu 20 bài: 14 bài đúng (70%), 6 bài sai (30%).
        """
        # 10 bài AI, 10 bài Con người
        # TP = 7, FN = 3 (Đúng 7/10 bài AI)
        # TN = 7, FP = 3 (Đúng 7/10 bài Người)
        # Accuracy = (7 + 7) / 20 = 70.0%
        y_true = [True] * 10 + [False] * 10
        y_pred = [True] * 7 + [False] * 3 + [False] * 7 + [True] * 3

        metrics = calculate_confusion_metrics(y_true, y_pred)

        self.assertEqual(metrics.total_samples, 20)
        self.assertEqual(metrics.tp, 7)
        self.assertEqual(metrics.tn, 7)
        self.assertEqual(metrics.fp, 3)
        self.assertEqual(metrics.fn, 3)
        self.assertAlmostEqual(metrics.accuracy, 0.70)
        self.assertAlmostEqual(metrics.precision, 0.70)
        self.assertAlmostEqual(metrics.recall, 0.70)
        self.assertAlmostEqual(metrics.specificity, 0.70)
        self.assertAlmostEqual(metrics.f1_score, 0.70)
        self.assertAlmostEqual(metrics.error_rate, 0.30)

    def test_zero_division_safety(self) -> None:
        """Kiểm tra độ bền khi tập mẫu rỗng hoặc không có dự đoán dương tính."""
        metrics = calculate_confusion_metrics([], [])
        self.assertEqual(metrics.accuracy, 0.0)
        self.assertEqual(metrics.precision, 0.0)
        self.assertEqual(metrics.recall, 0.0)
        self.assertEqual(metrics.f1_score, 0.0)

    def test_plot_confusion_matrix_generation(self) -> None:
        """Kiểm tra chức năng vẽ và lưu file ảnh Confusion Matrix."""
        # Tạo metrics mẫu 70%
        y_true = [True] * 10 + [False] * 10
        y_pred = [True] * 7 + [False] * 3 + [False] * 7 + [True] * 3
        metrics = calculate_confusion_metrics(y_true, y_pred)

        temp_img = ROOT_DIR / "data" / "outputs" / "test_cm_output.png"
        saved_path = plot_confusion_matrix(
            metrics=metrics,
            output_path=temp_img,
            title="Test Confusion Matrix (70% Accuracy Verification)",
            model_name="Unit Test Verification Engine",
        )

        self.assertTrue(saved_path.exists())
        self.assertGreater(os.path.getsize(saved_path), 1000)

        # Dọn dẹp sau khi kiểm tra
        if saved_path.exists():
            saved_path.unlink()

    def test_mock_evaluation_pipeline(self) -> None:
        """Kiểm tra chạy đánh giá pipeline với mock mode."""
        dummy_dataset = [
            {
                "media_id": f"id_{i}",
                "title": f"Song {i}",
                "artist": f"Artist {i}",
                "query": f"Song {i} Artist {i}",
                "gt_is_ai": i % 2 == 0,
                "gt_label": "AI" if i % 2 == 0 else "HUMAN",
                "row_idx": i,
            }
            for i in range(20)
        ]

        metrics, results = evaluate_dataset_with_pipeline(
            samples=dummy_dataset,
            mock_mode=True,
            mock_accuracy=0.70,
        )

        self.assertEqual(metrics.total_samples, 20)
        self.assertEqual(len(results), 20)
        self.assertGreaterEqual(metrics.accuracy, 0.0)
        self.assertLessEqual(metrics.accuracy, 1.0)

    def test_error_records_in_confusion_metrics(self) -> None:
        """Kiểm tra việc lưu vết error_records khi tính toán metrics."""
        err_recs = [
            {
                "Mã bài lỗi (Media ID)": "test_id",
                "Tên bài lỗi (Song Title)": "Error Song",
                "Ca sĩ (Artist)": "Test Artist",
                "Nhãn thực tế (Ground Truth)": "AI",
                "Nguyên nhân sự cố (Error)": "Zing VIP audio download error",
                "Bài nạp bù thay thế (Replacement)": "Backup Song",
                "Thời điểm ghi nhận": "2026-10-15 00:00:00",
            }
        ]
        metrics = calculate_confusion_metrics([True], [True], error_count=1, error_records=err_recs)
        self.assertEqual(metrics.errors, 1)
        self.assertEqual(len(metrics.error_records), 1)
        self.assertEqual(metrics.error_records[0]["Tên bài lỗi (Song Title)"], "Error Song")

    def test_export_excel_with_error_records_sheet(self) -> None:
        """Kiểm tra xuất file Excel với đầy đủ Sheet 3 'Sự cố & Nạp bù'."""
        err_recs = [
            {
                "Mã bài lỗi (Media ID)": "id_err",
                "Tên bài lỗi (Song Title)": "Err Track",
                "Ca sĩ (Artist)": "VIP Singer",
                "Nhãn thực tế (Ground Truth)": "AI",
                "Nguyên nhân sự cố (Error)": "403 Forbidden",
                "Bài nạp bù thay thế (Replacement)": "Backup Track",
                "Thời điểm ghi nhận": "2026-10-15 00:00:00",
            }
        ]
        metrics = calculate_confusion_metrics([True], [True], error_count=1, error_records=err_recs)
        results = [
            {
                "media_id": "id_1",
                "song_title": "Track 1",
                "artist": "Artist 1",
                "query": "Track 1 Artist 1",
                "ground_truth": "AI",
                "prediction": "AI",
                "is_correct": "ĐÚNG",
                "confidence_score": "95%",
                "reasons": "AI synthetic artifacts",
            }
        ]
        temp_excel = ROOT_DIR / "data" / "outputs" / "test_error_export.xlsx"
        try:
            exported_path = export_evaluation_excel(metrics, results, output_path=temp_excel)
            self.assertTrue(exported_path.exists())

            # Kiểm tra sheet names
            wb = openpyxl.load_workbook(str(exported_path), read_only=True)
            self.assertIn("Audit Details", wb.sheetnames)
            self.assertIn("Executive Summary", wb.sheetnames)
            self.assertIn("Sự cố & Nạp bù", wb.sheetnames)
            wb.close()
        finally:
            if temp_excel.exists():
                temp_excel.unlink()

    def test_load_metrics_from_audit_report_with_work_type_columns(self) -> None:
        """Kiểm tra đọc file audit có cả cột ground_truth và gt_work_type không bị nhận diện sai nhãn."""
        records = [
            {
                "media_id": "1",
                "ground_truth": "AI",
                "gt_work_type": "ORIGINAL",
                "pred_work_type": "ORIGINAL",
                "prediction": "AI",
            },
            {
                "media_id": "2",
                "ground_truth": "HUMAN",
                "gt_work_type": "COVER",
                "pred_work_type": "COVER",
                "prediction": "HUMAN",
            },
            {
                "media_id": "3",
                "ground_truth": "AI",
                "gt_work_type": "REMIX",
                "pred_work_type": "REMIX",
                "prediction": "HUMAN",
            },
            {
                "media_id": "4",
                "ground_truth": "HUMAN",
                "gt_work_type": "ORIGINAL",
                "pred_work_type": "ORIGINAL",
                "prediction": "AI",
            },
        ]
        temp_file = ROOT_DIR / "data" / "outputs" / "temp_test_audit.xlsx"
        try:
            with pd.ExcelWriter(temp_file, engine="openpyxl") as writer:
                pd.DataFrame(records).to_excel(writer, sheet_name="Audit Details", index=False)

            metrics, _ = load_metrics_from_audit_report(temp_file)
            self.assertEqual(metrics.total_samples, 4)
            self.assertEqual(metrics.tp, 1)  # id 1
            self.assertEqual(metrics.tn, 1)  # id 2
            self.assertEqual(metrics.fn, 1)  # id 3
            self.assertEqual(metrics.fp, 1)  # id 4
            self.assertAlmostEqual(metrics.accuracy, 0.5)
        finally:
            if temp_file.exists():
                temp_file.unlink()


# ==============================================================================
# HÀM CHÍNH (CLI RUNNER)
# ==============================================================================


def print_cli_report(metrics: ConfusionMetrics) -> None:
    """In báo cáo đánh giá trực quan trên terminal."""
    print("\n" + "=" * 76)
    print("        📊 BÁO CÁO ĐÁNH GIÁ ĐỘ CHÍNH XÁC & MA TRẬN NHẦM LẪN")
    print("=" * 76)
    print(f" Tổng số mẫu kiểm thử         : {metrics.total_samples}")
    print(f" Số mẫu thẩm định thành công   : {metrics.evaluated_samples}")
    print(f" Số mẫu bị sự cố kết nối      : {metrics.errors}")
    if metrics.error_records:
        print(
            f" 🔄 Số bài lỗi đã tự động nạp bù : {len(metrics.error_records)} bài (Chi tiết tại Sheet 'Sự cố & Nạp bù')"
        )
    print("-" * 76)
    print(" 🔢 MA TRẬN NHẦM LẪN (CONFUSION MATRIX 2x2):")
    print(
        f"   • True Positive  (TP - Đoán đúng AI)   : {metrics.tp} bài "
        f"({(metrics.tp / (metrics.tp + metrics.fn) * 100) if (metrics.tp + metrics.fn) > 0 else 0:.1f}% AI)"
    )
    print(
        f"   • True Negative  (TN - Đoán đúng Người): {metrics.tn} bài "
        f"({(metrics.tn / (metrics.tn + metrics.fp) * 100) if (metrics.tn + metrics.fp) > 0 else 0:.1f}% Người)"
    )
    print(
        f"   • False Positive (FP - Vu oan Người)   : {metrics.fp} bài "
        f"({(metrics.fp / (metrics.tn + metrics.fp) * 100) if (metrics.tn + metrics.fp) > 0 else 0:.1f}% Báo động giả)"
    )
    print(
        f"   • False Negative (FN - Bỏ sót AI)      : {metrics.fn} bài "
        f"({(metrics.fn / (metrics.tp + metrics.fn) * 100) if (metrics.tp + metrics.fn) > 0 else 0:.1f}% Lọt lưới)"
    )
    print("-" * 76)
    print(" 🎯 CHỈ SỐ ĐO LƯỜNG:")
    print(f"   ⭐ Accuracy (Độ chính xác tổng thể)   : {metrics.accuracy * 100:.2f}%")
    print(f"   🎯 Precision (Độ tin cậy khi báo AI)   : {metrics.precision * 100:.2f}%")
    print(f"   📡 Recall (Độ nhạy phát hiện AI)       : {metrics.recall * 100:.2f}%")
    print(f"   🛡️ Specificity (Độ đặc hiệu nhạc người): {metrics.specificity * 100:.2f}%")
    print(f"   ⚖️ F1-Score (Cân bằng Precision/Recall): {metrics.f1_score * 100:.2f}%")
    print(f"   ⚖️ Balanced Accuracy                   : {metrics.balanced_accuracy * 100:.2f}%")
    print(f"   ⚠️ Error Rate (Tỷ lệ sai sót)          : {metrics.error_rate * 100:.2f}%")
    print("=" * 76)


def main() -> None:
    """Điểm vào thực thi dòng lệnh của công cụ."""
    parser = argparse.ArgumentParser(
        description="Đánh giá độ chính xác & Vẽ Ma trận nhầm lẫn (Confusion Matrix) cho AI Music Detector",
    )
    parser.add_argument(
        "--file",
        default="testcase.xlsx",
        help="File Excel ground truth (mặc định: testcase.xlsx)",
    )
    parser.add_argument(
        "--sheet",
        default="Media_W38",
        help="Sheet Excel cần thẩm định (mặc định: Media_W38)",
    )
    parser.add_argument(
        "--sample",
        type=int,
        default=20,
        help="Số lượng bài mẫu thẩm định (mặc định: 20)",
    )
    parser.add_argument(
        "--balanced",
        action="store_true",
        default=True,
        help="Lấy mẫu cân bằng 50/50 giữa AI và Con người",
    )
    parser.add_argument(
        "--all",
        action="store_true",
        help="Thẩm định toàn bộ dataset",
    )
    parser.add_argument(
        "--from-report",
        type=str,
        default=None,
        help="Đọc kết quả từ file Excel audit trước đó để vẽ Confusion Matrix trực tiếp",
    )
    parser.add_argument(
        "--mock",
        action="store_true",
        help="Chạy ở chế độ mô phỏng (offline/mock) để kiểm tra nhanh mà không gọi API",
    )
    parser.add_argument(
        "--mock-acc",
        type=float,
        default=0.70,
        help="Tỉ lệ chính xác giả lập khi chạy --mock (mặc định: 0.70)",
    )
    parser.add_argument(
        "--output",
        type=str,
        default=None,
        help="Đường dẫn lưu hình ảnh Confusion Matrix",
    )
    parser.add_argument(
        "--unit-test",
        action="store_true",
        help="Chạy toàn bộ bộ kiểm thử tự động unittest",
    )

    args = parser.parse_args()

    if args.unit_test:
        print("🚀 Đang chạy bộ kiểm thử tự động (Unit Test)...")
        suite = unittest.TestLoader().loadTestsFromTestCase(TestConfusionMatrixEvaluation)
        runner = unittest.TextTestRunner(verbosity=2)
        result = runner.run(suite)
        sys.exit(0 if result.wasSuccessful() else 1)

    if args.from_report:
        print(f"📖 Đang đọc kết quả đối soát từ: {args.from_report}")
        metrics, _ = load_metrics_from_audit_report(args.from_report)
        print_cli_report(metrics)
        saved_img = plot_confusion_matrix(
            metrics=metrics,
            output_path=args.output,
            title="Confusion Matrix (Từ Báo Cáo Đối Soát)",
            model_name="Imported Audit Report",
        )
        print(f"\n🖼️  Đã xuất hình ảnh Confusion Matrix tại: {saved_img}\n")
        return

    # Chế độ chạy benchmark
    dataset_file = ROOT_DIR / args.file
    if not dataset_file.exists():
        print(f"⚠️  Không tìm thấy file '{args.file}'. Chuyển sang chế độ chạy Unit Test...")
        unittest.main(argv=["first-arg-is-ignored"], exit=False)
        return

    print(f"📖 Đang đọc dataset từ '{args.file}' [Sheet: '{args.sheet}']...")
    dataset = load_testcase_dataset(file_path=args.file, sheet_name=args.sheet)

    if args.all:
        samples = dataset
    else:
        samples = select_samples(dataset, sample_size=args.sample, balanced=args.balanced)

    sample_ai = sum(1 for s in samples if s["gt_is_ai"])
    sample_human = sum(1 for s in samples if not s["gt_is_ai"])
    print(f"🎯 Lấy mẫu: {len(samples)} bài (AI: {sample_ai} bài | Con người: {sample_human} bài)")

    if args.mock:
        print(f"⚡ Đang chạy chế độ MOCK (Mô phỏng tỉ lệ {args.mock_acc * 100:.1f}%)...")
    else:
        print("🚀 Đang thẩm định qua hệ thống AI Music Detector...")

    metrics, results = evaluate_dataset_with_pipeline(
        samples=samples,
        dataset=dataset,
        mock_mode=args.mock,
        mock_accuracy=args.mock_acc,
    )

    print_cli_report(metrics)
    saved_img = plot_confusion_matrix(
        metrics=metrics,
        output_path=args.output,
        title=f"AI Music Detector - Confusion Matrix ({metrics.accuracy * 100:.1f}% Accuracy)",
        model_name="Mock Engine" if args.mock else "Project Core AI Engine",
    )
    print(f"\n🖼️  Đã xuất hình ảnh Confusion Matrix tại: {saved_img}")
    audit_file = export_evaluation_excel(metrics, results)
    print(f"📁 Đã xuất bảng đối soát Excel chi tiết tại: {audit_file}\n")


if __name__ == "__main__":
    main()
