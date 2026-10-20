"""
Công cụ Đánh Giá & Chấm Điểm Độ Chính Xác (Benchmark & Accuracy Evaluator)
Sử dụng dữ liệu Ground Truth từ file testcase.xlsx (cột 'AI Yes').

Các chỉ số đo lường:
- Confusion Matrix: True Positive (TP), True Negative (TN), False Positive (FP), False Negative (FN)
- Accuracy (Độ chính xác tổng thể)
- Precision (Độ tin cậy khi gắn nhãn AI - tránh bắt oan)
- Recall / Sensitivity (Độ nhạy bắt nhạc AI - tránh lọt lưới)
- Specificity (Độ đặc hiệu - khả năng nhận diện đúng nhạc con người)
- F1-Score (Điểm điều hòa cân bằng giữa Precision và Recall)

Cách dùng:
  python benchmark_evaluator.py --sample 20 --balanced
  python benchmark_evaluator.py --sample 50
  python benchmark_evaluator.py --all
"""

import argparse
import os
import random
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any

# Thêm root vào sys.path
root_dir = Path(__file__).resolve().parent
if str(root_dir) not in sys.path:
    sys.path.insert(0, str(root_dir))

if sys.platform == "win32":
    try:
        if hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(encoding="utf-8")
        if hasattr(sys.stderr, "reconfigure"):
            sys.stderr.reconfigure(encoding="utf-8")
    except Exception:  # noqa: BLE001, S110
        pass

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd

from backend.core.container import get_container
from backend.core.models.entities import (
    DetectionResult,
    TrackItem,
)
from backend.services.work_type_resolver import resolve_work_type_from_metadata


def map_song_type_to_work_type(song_type_val: str | None) -> str:
    """Ánh xạ nhãn 'Song Type' từ testcase sang chuẩn WorkType enum value."""
    st_clean = (str(song_type_val) if song_type_val is not None else "").strip().lower()
    if not st_clean or st_clean in ("none", "nan", "null"):
        return "UNCERTAIN"
    if "cover" in st_clean:
        return "COVER"
    if any(k in st_clean for k in ["remix", "remake"]):
        return "REMIX"
    if any(k in st_clean for k in ["original", "live"]):
        return "ORIGINAL"
    return "UNCERTAIN"


def calculate_work_type_confusion_matrix(
    y_true: list[str],
    y_pred: list[str],
    labels: list[str] | None = None,
) -> dict[str, Any]:
    """Tính toán ma trận nhầm lẫn đa lớp và chỉ số hiệu năng cho Work Type."""
    if labels is None:
        labels = ["ORIGINAL", "COVER", "REMIX"]

    all_pred_cols = list(labels)
    if "UNCERTAIN" not in all_pred_cols:
        all_pred_cols.append("UNCERTAIN")

    matrix: dict[str, dict[str, int]] = {gt: {pred: 0 for pred in all_pred_cols} for gt in labels}

    total = 0
    correct = 0

    for yt, yp in zip(y_true, y_pred, strict=False):
        yt_norm = str(yt or "").upper().strip()
        yp_norm = str(yp or "").upper().strip()
        if yt_norm in matrix:
            if yp_norm not in matrix[yt_norm]:
                yp_norm = "UNCERTAIN"
            matrix[yt_norm][yp_norm] += 1
            total += 1
            if yt_norm == yp_norm:
                correct += 1

    accuracy = (correct / total) if total > 0 else 0.0

    per_class: dict[str, dict[str, Any]] = {}
    for label in labels:
        tp = matrix.get(label, {}).get(label, 0)
        fn = sum(matrix.get(label, {}).get(p, 0) for p in all_pred_cols if p != label)
        fp = sum(matrix.get(other, {}).get(label, 0) for other in labels if other != label)

        prec = tp / (tp + fp) if (tp + fp) > 0 else 0.0
        rec = tp / (tp + fn) if (tp + fn) > 0 else 0.0
        f1 = (2 * prec * rec / (prec + rec)) if (prec + rec) > 0 else 0.0

        per_class[label] = {
            "total": tp + fn,
            "tp": tp,
            "fp": fp,
            "fn": fn,
            "precision": prec,
            "recall": rec,
            "f1_score": f1,
        }

    return {
        "matrix": matrix,
        "labels": labels,
        "total_evaluated": total,
        "correct": correct,
        "accuracy": accuracy,
        "per_class": per_class,
    }


def plot_work_type_confusion_matrix(
    metrics: dict[str, Any],
    output_path: str | Path | None = None,
    title: str = "Ma Trận Nhầm Lẫn Tác Phẩm (Original vs Cover vs Remix)",
) -> Path:
    """Vẽ ma trận nhầm lẫn đa lớp cho Work Type và lưu file ảnh PNG (300 DPI)."""
    if output_path is None:
        out_dir = root_dir / "data" / "outputs"
        out_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        target_file = out_dir / f"work_type_cm_{ts}.png"
    else:
        target_file = Path(output_path)
        target_file.parent.mkdir(parents=True, exist_ok=True)

    labels = metrics.get("labels", ["ORIGINAL", "COVER", "REMIX"])
    matrix_dict = metrics.get("matrix", {})

    matrix_data = []
    for gt in labels:
        row = [matrix_dict.get(gt, {}).get(pred, 0) for pred in labels]
        matrix_data.append(row)
    matrix_arr = np.array(matrix_data)

    fig, (ax_cm, ax_stats) = plt.subplots(
        nrows=1,
        ncols=2,
        figsize=(15, 6.8),
        dpi=300,
        gridspec_kw={"width_ratios": [1.2, 1.0]},
    )
    fig.patch.set_facecolor("#FAFAFC")

    # 1. Heatmap bên trái
    cmap = plt.cm.Purples
    im = ax_cm.imshow(matrix_arr, interpolation="nearest", cmap=cmap)
    plt.colorbar(im, ax=ax_cm, fraction=0.046, pad=0.04, label="Số lượng bài hát")

    tick_marks = np.arange(len(labels))
    ax_cm.set_xticks(tick_marks)
    ax_cm.set_xticklabels([f"Đoán: {lbl}" for lbl in labels], fontsize=10, fontweight="bold")
    ax_cm.set_yticks(tick_marks)
    ax_cm.set_yticklabels([f"Thực tế: {lbl}" for lbl in labels], fontsize=10, fontweight="bold")
    ax_cm.set_xlabel("DỰ ĐOÁN CỦA HỆ THỐNG (PREDICTED)", fontsize=11, labelpad=10)
    ax_cm.set_ylabel("NHÃN THỰC TẾ (GROUND TRUTH)", fontsize=11, labelpad=10)
    ax_cm.set_title(title, fontsize=12, fontweight="bold", pad=12)

    row_sums = matrix_arr.sum(axis=1)
    max_val = matrix_arr.max() if matrix_arr.size > 0 else 1
    for i in range(len(labels)):
        for j in range(len(labels)):
            val = matrix_arr[i, j]
            total_row = row_sums[i]
            pct = (val / total_row * 100) if total_row > 0 else 0.0
            color = "white" if val > (max_val / 2) else "black"
            ax_cm.text(
                j,
                i,
                f"{val}\n({pct:.1f}%)",
                ha="center",
                va="center",
                color=color,
                fontsize=11,
                fontweight="bold",
            )

    # 2. Bảng thống kê bên phải
    ax_stats.axis("off")
    acc = metrics.get("accuracy", 0.0)
    tot = metrics.get("total_evaluated", 0)
    cor = metrics.get("correct", 0)

    header_text = (
        f"[THỐNG KÊ HIỆU NĂNG PHÂN BIỆT TÁC PHẨM]\n"
        f"• Tổng số bài đánh giá : {tot} bài\n"
        f"• Số bài đoán đúng    : {cor} bài\n"
        f"• Độ chính xác chung   : {acc * 100:.2f}%\n"
        f"────────────────────────────────────────────────"
    )
    ax_stats.text(
        0.05, 0.95, header_text, fontsize=11, family="sans-serif", va="top", fontweight="bold"
    )

    y_pos = 0.65
    per_class = metrics.get("per_class", {})
    for lbl in labels:
        c_met = per_class.get(lbl, {})
        c_tot = c_met.get("total", 0)
        c_prec = c_met.get("precision", 0.0) * 100
        c_rec = c_met.get("recall", 0.0) * 100
        c_f1 = c_met.get("f1_score", 0.0) * 100

        class_text = (
            f"[NHÃN]: {lbl} (Mẫu thực tế: {c_tot} bài)\n"
            f"   - Precision (Độ chuẩn xác) : {c_prec:.1f}%\n"
            f"   - Recall (Độ bao phủ)      : {c_rec:.1f}%\n"
            f"   - F1-Score (Cân bằng)      : {c_f1:.1f}%\n"
        )
        ax_stats.text(0.05, y_pos, class_text, fontsize=10, family="sans-serif", va="top")
        y_pos -= 0.18

    plt.tight_layout()
    fig.savefig(target_file, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return target_file


def load_testcase_dataset(
    file_path: str = "testcase.xlsx", sheet_name: str = "Media_W38"
) -> list[dict[str, Any]]:
    """Đọc dữ liệu từ file Excel testcase và trích xuất danh sách bài hát kèm nhãn Ground Truth."""
    if not os.path.exists(file_path):
        data_candidate = os.path.join(root_dir, "data", file_path)
        if os.path.exists(data_candidate):
            file_path = data_candidate
        else:
            raise FileNotFoundError(f"Không tìm thấy file: {file_path}")

    wb = openpyxl.load_workbook(file_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        sheet_name = wb.sheetnames[0]

    sheet = wb[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Sheet '{sheet_name}' không có dữ liệu.")

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]

    # Tìm các cột cần thiết
    title_idx = -1
    artist_idx = -1
    ai_col_idx = -1
    media_id_idx = -1
    note_idx = -1
    song_type_idx = -1
    release_date_idx = -1
    distributor_idx = -1

    for idx, h in enumerate(headers):
        h_low = h.lower()
        if "song title" in h_low or "tên bài" in h_low or "title" in h_low:
            title_idx = idx
        elif "artist" in h_low or "ca sĩ" in h_low or "nghệ sĩ" in h_low:
            artist_idx = idx
        elif "ai yes" in h_low or h_low == "ai":
            ai_col_idx = idx
        elif "media_id" in h_low or h_low == "id":
            media_id_idx = idx
        elif "note" in h_low or "ghi chú" in h_low:
            note_idx = idx
        elif "song type" in h_low or "loại bài" in h_low or "type" in h_low:
            song_type_idx = idx
        elif "release" in h_low or "ngày phát hành" in h_low:
            release_date_idx = idx
        elif "distributor" in h_low or "nhà phân phối" in h_low:
            distributor_idx = idx

    if title_idx == -1:
        title_idx = 1 if len(headers) > 1 else 0

    dataset: list[dict[str, Any]] = []
    for r_idx, r in enumerate(rows[1:], start=2):
        title = str(r[title_idx]).strip() if r[title_idx] is not None else ""
        if not title:
            continue

        artist = (
            str(r[artist_idx]).strip() if (artist_idx != -1 and r[artist_idx] is not None) else ""
        )
        media_id = (
            str(r[media_id_idx]).strip()
            if (media_id_idx != -1 and r[media_id_idx] is not None)
            else f"row_{r_idx}"
        )
        note = str(r[note_idx]).strip() if (note_idx != -1 and r[note_idx] is not None) else ""
        song_type = (
            str(r[song_type_idx]).strip()
            if (song_type_idx != -1 and r[song_type_idx] is not None)
            else ""
        )
        release_date = (
            str(r[release_date_idx]).strip()
            if (release_date_idx != -1 and r[release_date_idx] is not None)
            else ""
        )

        distributor = (
            str(r[distributor_idx]).strip()
            if (distributor_idx != -1 and r[distributor_idx] is not None)
            else ""
        )

        ai_val = (
            str(r[ai_col_idx]).strip().lower()
            if (ai_col_idx != -1 and r[ai_col_idx] is not None)
            else ""
        )
        # Nhãn thực tế (Ground Truth)
        is_ai_ground_truth = "yes" in ai_val or ai_val == "y"
        gt_work_type = map_song_type_to_work_type(song_type)

        query = f"{title} {artist}".strip()
        dataset.append(
            {
                "media_id": media_id,
                "title": title,
                "artist": artist,
                "query": query,
                "gt_is_ai": is_ai_ground_truth,
                "gt_label": "AI" if is_ai_ground_truth else "HUMAN",
                "gt_work_type": gt_work_type,
                "note": note,
                "song_type": song_type,
                "release_date": release_date,
                "distributor": distributor,
                "row_idx": r_idx,
            }
        )

    return dataset


def select_samples(
    dataset: list[dict[str, Any]],
    sample_size: int | None = 20,
    balanced: bool = True,
    seed: int = 42,
) -> list[dict[str, Any]]:
    """Lấy mẫu bài hát để thẩm định (hỗ trợ chế độ cân bằng phân tầng 50% AI - 50% Human)."""
    if sample_size is None or sample_size >= len(dataset):
        return dataset

    random.seed(seed)
    ai_items = [d for d in dataset if d["gt_is_ai"]]
    human_items = [d for d in dataset if not d["gt_is_ai"]]

    if balanced:
        half = sample_size // 2
        selected_ai = random.sample(ai_items, min(half, len(ai_items)))
        remaining = sample_size - len(selected_ai)
        selected_human = random.sample(human_items, min(remaining, len(human_items)))
        combined = selected_ai + selected_human
        random.shuffle(combined)
        return combined
    else:
        return random.sample(dataset, sample_size)


def run_benchmark(
    samples: list[dict[str, Any]],
    dataset: list[dict[str, Any]] | None = None,
    model_override: str | None = None,
    force_rescan: bool = False,
) -> dict[str, Any]:
    """
    Thực hiện thẩm định toàn bộ mẫu và tính toán các chỉ số độ chính xác.
    Hỗ trợ cơ chế Nạp Bù Mẫu Tự Động (Resilient Sampling Pool):
    Khi 1 bài hát gặp lỗi tải/Zing VIP, tự động lấy bài kế tiếp cùng nhãn từ dataset dự phòng
    để đảm bảo đủ chỉ tiêu số bài thành công, đồng thời lưu vết toàn bộ bài lỗi.
    """
    container = get_container()
    detection_service = container.detection_service
    resolver = container.resolver

    results: list[dict[str, Any]] = []
    error_records: list[dict[str, Any]] = []
    y_true_work: list[str] = []
    y_pred_work: list[str] = []

    tp = 0  # Thực tế AI, Đoán AI (True Positive)
    tn = 0  # Thực tế Người, Đoán Người (True Negative)
    fp = 0  # Thực tế Người, Đoán AI (False Positive - Vu oan)
    fn = 0  # Thực tế AI, Đoán Người (False Negative - Bỏ sót)
    errors = 0
    unverified_matches = 0  # Số bài không tìm thấy nguồn audio khớp trên YouTube/Web

    evaluation_queue = list(samples)
    attempted_ids = {s["media_id"] for s in samples}

    print("=" * 75)
    print(f"🚀 BẮT ĐẦU CHẤM ĐIỂM TESTCASE (MỤC TIÊU: {len(samples)} BÀI HÁT THÀNH CÔNG)")
    print("=" * 75)

    i = 0
    while i < len(evaluation_queue):
        item = evaluation_queue[i]
        i += 1
        gt = item["gt_label"]
        query = item["query"]
        print(f"\n[{i}/{len(evaluation_queue)}] Thẩm định: '{query}' | Ground Truth: {gt}")

        track: TrackItem | None = None
        try:
            try:
                track = resolver.resolve_track(
                    item["title"], artist=item.get("artist", ""), force_rescan=force_rescan
                )
            except TypeError:
                track = resolver.resolve_track(item["title"], artist=item.get("artist", ""))

            track.song_type = item.get("song_type") or ""
            track.release_date = str(item.get("release_date") or "")
            track.distributor = str(item.get("distributor") or "")
            try:
                res: DetectionResult = detection_service.detect_single(
                    track, force_rescan=force_rescan, model_override=model_override
                )
            except TypeError:
                res = detection_service.detect_single(track, model_override=model_override)

            # Đọc lại trạng thái xác thực nguồn sau khi audio_engine đã chạy và xử lý fallback
            is_verified = track.is_verified_match is not False
            if not is_verified:
                unverified_matches += 1
                print(
                    f"       -> ⚠️ CẢNH BÁO NGUỒN: YouTube không tìm thấy video khớp! (Score: {track.match_score})"
                )

            pred_is_ai = res.is_ai
            pred_label = "AI" if pred_is_ai else "HUMAN"
            is_correct = pred_is_ai == item["gt_is_ai"]

            # Đánh giá Work Type (Original vs Cover vs Remix)
            gt_work = item.get("gt_work_type", "UNCERTAIN")
            raw_w_type = getattr(res, "work_type", None)
            pred_work = (
                raw_w_type.value
                if hasattr(raw_w_type, "value")
                else (str(raw_w_type) if raw_w_type else "UNCERTAIN")
            )
            # Tự động giải mã tác phẩm nếu kết quả từ model hoặc cache DB cũ vẫn là UNCERTAIN
            if pred_work in ("UNCERTAIN", "ERROR", "", "NAN"):
                resolved_w, resolved_label, is_cov, orig_a, orig_t = (
                    resolve_work_type_from_metadata(
                        title=item.get("title") or track.track_title or track.raw_query,
                        artist=item.get("artist") or track.artist or "",
                        metadata=item,
                    )
                )
                pred_work = resolved_w.value
                res.work_type = resolved_w
                res.work_label = resolved_label
                res.is_cover = is_cov
                if orig_a and not getattr(res, "original_artist", None):
                    res.original_artist = orig_a
                if orig_t and not getattr(res, "original_song_title", None):
                    res.original_song_title = orig_t

            work_is_correct = (pred_work == gt_work) if gt_work != "UNCERTAIN" else None
            if gt_work != "UNCERTAIN":
                y_true_work.append(gt_work)
                y_pred_work.append(pred_work)

            # Phân loại Confusion Matrix
            if item["gt_is_ai"] and pred_is_ai:
                metric_type = "TP (Đúng: Phát hiện AI)"
                tp += 1
            elif not item["gt_is_ai"] and not pred_is_ai:
                metric_type = "TN (Đúng: Xác nhận Người)"
                tn += 1
            elif not item["gt_is_ai"] and pred_is_ai:
                metric_type = "FP (Sai: Vu oan nhạc người)"
                fp += 1
            else:
                metric_type = "FN (Sai: Bỏ sót nhạc AI)"
                fn += 1

            icon = "✅" if is_correct else "❌"
            match_status_text = "ĐÃ XÁC THỰC" if is_verified else "SAI LỆCH (UNVERIFIED)"
            ai_model_is_ai = getattr(res, "ai_model_is_ai", res.is_ai)
            ai_pred_label = "AI" if ai_model_is_ai else "HUMAN"
            ai_verdict_str = res.ai_model_verdict or ai_pred_label
            ai_is_correct = "ĐÚNG" if ai_model_is_ai == item["gt_is_ai"] else "SAI"

            consensus_str = res.consensus_verdict or (
                "AI" if getattr(res, "consensus_is_ai", res.is_ai) else "HUMAN"
            )

            work_icon = "💿" if pred_work == "ORIGINAL" else ("🎙️" if pred_work == "COVER" else "🎛️")
            print(
                f"       -> Model AI: {ai_verdict_str} | Tổng hợp: {consensus_str} ({res.verdict_label}) | {icon} {metric_type}"
            )
            print(
                f"       -> Tác phẩm: {work_icon} {getattr(res, 'work_label', pred_work)} (Thực tế: {gt_work})"
            )
            print(
                f"       -> Nguồn audio: {match_status_text} | Độ tin cậy: {res.confidence_score * 100:.1f}%"
            )
            print(f"       -> Lý do: {res.reasons[0] if res.reasons else 'N/A'}")

            results.append(
                {
                    "media_id": item["media_id"],
                    "song_title": item["title"],
                    "artist": item["artist"],
                    "query": query,
                    "ground_truth": gt,
                    "gt_work_type": gt_work,
                    "pred_work_type": pred_work,
                    "work_is_correct": "ĐÚNG"
                    if work_is_correct is True
                    else ("SAI" if work_is_correct is False else "N/A"),
                    "work_label": getattr(res, "work_label", "Chưa rõ tác phẩm"),
                    "original_artist": getattr(res, "original_artist", "") or "N/A",
                    "original_song_title": getattr(res, "original_song_title", "") or "N/A",
                    "is_cover": getattr(res, "is_cover", False),
                    "ai_model_prediction": ai_pred_label,
                    "ai_model_verdict": ai_verdict_str,
                    "ai_model_is_correct": ai_is_correct,
                    "consensus_prediction": pred_label,
                    "consensus_verdict": consensus_str,
                    "prediction": pred_label,
                    "is_correct": "ĐÚNG" if is_correct else "SAI",
                    "classification_type": metric_type,
                    "is_verified_source": is_verified,
                    "source_verification": match_status_text,
                    "match_score": (
                        f"{track.match_score * 100:.1f}%"
                        if track.match_score is not None
                        else "N/A"
                    ),
                    "verdict_type": (
                        res.verdict_type.value
                        if hasattr(res.verdict_type, "value")
                        else str(res.verdict_type)
                    ),
                    "verdict_label": res.verdict_label,
                    "confidence_score": f"{res.confidence_score * 100:.1f}%",
                    "suspected_engine": res.suspected_engine,
                    "source_url": (res.source_url or getattr(track, "resolved_url", "") or ""),
                    "reasons": " | ".join(res.reasons or []),
                    "error": "",
                }
            )

        except Exception as e:  # noqa: BLE001
            errors += 1
            print(f"       -> ❌ LỖI XỬ LÝ NGUỒN: {e}")

            # Tìm mẫu thay thế cùng nhãn từ dataset dự phòng
            replacement: dict[str, Any] | None = None
            if dataset:
                for cand in dataset:
                    if (
                        cand["gt_is_ai"] == item["gt_is_ai"]
                        and cand["media_id"] not in attempted_ids
                    ):
                        replacement = cand
                        attempted_ids.add(cand["media_id"])
                        evaluation_queue.append(cand)
                        break

            # Ghi nhận sự cố vào error_records
            error_records.append(
                {
                    "Mã bài lỗi (Media ID)": item["media_id"],
                    "Tên bài lỗi (Song Title)": item["title"],
                    "Ca sĩ (Artist)": item["artist"],
                    "Nhãn thực tế (Ground Truth)": gt,
                    "Nguyên nhân sự cố (Error)": str(e),
                    "Bài nạp bù thay thế (Replacement)": replacement["query"]
                    if replacement
                    else "Không còn bài dự phòng",
                    "Thời điểm ghi nhận": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }
            )

            if replacement:
                print(
                    f"       -> 🔄 ĐÃ NẠP BÙ MẪU THAY THẾ: '{replacement['query']}' (cùng nhãn {gt})"
                )
            else:
                # Nếu không tìm thấy bài nạp bù, ghi nhận là kết quả ERROR
                results.append(
                    {
                        "media_id": item["media_id"],
                        "song_title": item["title"],
                        "artist": item["artist"],
                        "query": query,
                        "ground_truth": gt,
                        "ai_model_prediction": "ERROR",
                        "ai_model_verdict": "ERROR",
                        "ai_model_is_correct": "LỖI",
                        "consensus_prediction": "ERROR",
                        "consensus_verdict": "ERROR",
                        "prediction": "ERROR",
                        "is_correct": "LỖI",
                        "classification_type": "ERROR",
                        "is_verified_source": False,
                        "source_verification": "LỖI",
                        "match_score": "0%",
                        "verdict_type": "ERROR",
                        "verdict_label": "LỖI HỆ THỐNG",
                        "confidence_score": "0%",
                        "suspected_engine": "N/A",
                        "source_url": (
                            getattr(track, "resolved_url", "")
                            if track and getattr(track, "resolved_url", None)
                            else ""
                        ),
                        "reasons": str(e),
                        "error": str(e),
                    }
                )

        if i < len(evaluation_queue):
            time.sleep(2.5)

    # Tính toán các chỉ số
    evaluated_total = tp + tn + fp + fn
    accuracy = ((tp + tn) / evaluated_total) if evaluated_total > 0 else 0.0
    precision = (tp / (tp + fp)) if (tp + fp) > 0 else 0.0
    recall = (tp / (tp + fn)) if (tp + fn) > 0 else 0.0
    specificity = (tn / (tn + fp)) if (tn + fp) > 0 else 0.0
    f1_score = (2 * precision * recall / (precision + recall)) if (precision + recall) > 0 else 0.0

    # Đánh giá chỉ số phân biệt Tác phẩm (Work Type: Original vs Cover vs Remix)
    work_type_metrics = calculate_work_type_confusion_matrix(y_true_work, y_pred_work)
    work_type_cm_img_path: str | None = None
    if work_type_metrics.get("total_evaluated", 0) > 0:
        try:
            img_file = plot_work_type_confusion_matrix(work_type_metrics)
            work_type_cm_img_path = str(img_file)
        except Exception as plot_err:  # noqa: BLE001
            print(f"⚠️ Không thể tạo ảnh ma trận nhầm lẫn tác phẩm: {plot_err}")

    # Tính riêng chỉ số cho tập các bài tìm ĐÚNG nguồn âm thanh (Strictly Verified Tracks)
    verified_results = [r for r in results if r.get("is_verified_source")]
    v_tp = sum(1 for r in verified_results if r["ground_truth"] == "AI" and r["prediction"] == "AI")
    v_tn = sum(
        1 for r in verified_results if r["ground_truth"] == "HUMAN" and r["prediction"] == "HUMAN"
    )
    v_fp = sum(
        1 for r in verified_results if r["ground_truth"] == "HUMAN" and r["prediction"] == "AI"
    )
    v_fn = sum(
        1 for r in verified_results if r["ground_truth"] == "AI" and r["prediction"] == "HUMAN"
    )
    v_total = v_tp + v_tn + v_fp + v_fn
    v_acc = ((v_tp + v_tn) / v_total) if v_total > 0 else 0.0
    v_rec = (v_tp / (v_tp + v_fn)) if (v_tp + v_fn) > 0 else 0.0
    v_prec = (v_tp / (v_tp + v_fp)) if (v_tp + v_fp) > 0 else 0.0
    v_f1 = (2 * v_prec * v_rec / (v_prec + v_rec)) if (v_prec + v_rec) > 0 else 0.0

    metrics_summary = {
        "total_samples": len(samples),
        "evaluated_successfully": evaluated_total,
        "errors": errors,
        "unverified_matches": unverified_matches,
        "TP": tp,
        "TN": tn,
        "FP": fp,
        "FN": fn,
        "accuracy": accuracy,
        "precision": precision,
        "recall": recall,
        "specificity": specificity,
        "f1_score": f1_score,
        "error_records": error_records,
        "replaced_errors_count": len(error_records),
        "work_type_metrics": work_type_metrics,
        "work_type_cm_image_path": work_type_cm_img_path,
        "verified_metrics": {
            "verified_samples": v_total,
            "TP": v_tp,
            "TN": v_tn,
            "FP": v_fp,
            "FN": v_fn,
            "accuracy": v_acc,
            "recall": v_rec,
            "precision": v_prec,
            "f1_score": v_f1,
        },
    }

    return {
        "metrics": metrics_summary,
        "results": results,
        "unverified_matches": unverified_matches,
    }


def print_summary_report(metrics: dict[str, Any], results: list[dict[str, Any]]):
    """In bảng báo cáo tổng kết chi tiết ra màn hình console."""
    print("\n" + "=" * 75)
    print("                📊 BẢNG TỔNG KẾT ĐÁNH GIÁ ĐỘ CHÍNH XÁC")
    print("=" * 75)
    print(f" Tổng số bài kiểm tra    : {metrics['total_samples']}")
    print(f" Số bài thành công        : {metrics['evaluated_successfully']}")
    print(f" Số bài bị lỗi kết nối    : {metrics['errors']}")
    if metrics.get("replaced_errors_count", 0) > 0:
        print(
            f" 🔄 Số bài lỗi đã tự động nạp bù : {metrics['replaced_errors_count']} bài (Chi tiết tại Sheet 'Sự cố & Nạp bù')"
        )
    print("-" * 75)
    print(" 🔢 MA TRẬN NHẦM LẪN AI/HUMAN (CONFUSION MATRIX):")
    print(f"   • True Positive  (TP - Đoán đúng AI)   : {metrics['TP']} bài")
    print(f"   • True Negative  (TN - Đoán đúng Người): {metrics['TN']} bài")
    print(f"   • False Positive (FP - Vu oan Người)   : {metrics['FP']} bài (⚠️ Báo động giả)")
    print(f"   • False Negative (FN - Bỏ sót AI)      : {metrics['FN']} bài (⚠️ Lọt lưới)")
    print("-" * 75)
    print(" 🎯 CÁC CHỈ SỐ ĐO LƯỜNG PHÂN BIỆT AI:")
    print(f"   ⭐ Accuracy (Độ chính xác tổng)        : {metrics['accuracy'] * 100:.2f}%")
    print(f"   🎯 Precision (Độ tin cậy khi báo AI)   : {metrics['precision'] * 100:.2f}%")
    print(f"   📡 Recall (Độ nhạy bắt bài AI)         : {metrics['recall'] * 100:.2f}%")
    print(f"   🛡️ Specificity (Độ đặc hiệu nhạc người): {metrics['specificity'] * 100:.2f}%")
    print(f"   ⚖️ F1-Score (Cân bằng Precision/Recall): {metrics['f1_score'] * 100:.2f}%")
    print("=" * 75)

    wt_metrics = metrics.get("work_type_metrics")
    if wt_metrics and wt_metrics.get("total_evaluated", 0) > 0:
        print("\n" + "=" * 75)
        print("      💿 BẢNG ĐỐI SOÁT PHÂN BIỆT TÁC PHẨM (ORIGINAL vs COVER vs REMIX)")
        print("=" * 75)
        print(f" Tổng số bài có nhãn tác phẩm : {wt_metrics['total_evaluated']} bài")
        print(f" Số bài nhận diện đúng tác phẩm: {wt_metrics['correct']} bài")
        print(f" ⭐ Độ chính xác tác phẩm chung : {wt_metrics['accuracy'] * 100:.2f}%")
        print("-" * 75)
        per_class = wt_metrics.get("per_class", {})
        for lbl, c_met in per_class.items():
            print(
                f"   🏷️ {lbl:<8}: Mẫu={c_met['total']:<2} | "
                f"Precision={c_met['precision'] * 100:.1f}% | "
                f"Recall={c_met['recall'] * 100:.1f}% | "
                f"F1-Score={c_met['f1_score'] * 100:.1f}%"
            )
        if metrics.get("work_type_cm_image_path"):
            print(f" 🖼️ Ma trận nhầm lẫn tác phẩm đã lưu: {metrics['work_type_cm_image_path']}")
        print("=" * 75)

    if metrics.get("unverified_matches", 0) > 0:
        print("\n" + "!" * 75)
        print(
            f" ⚠️ CẢNH BÁO NGUỒN ÂM THANH: Có {metrics['unverified_matches']} bài không tìm thấy trên YouTube (tải video lạc đề)!"
        )
        vm = metrics.get("verified_metrics", {})
        if vm and vm.get("verified_samples", 0) > 0:
            print(f" 🎯 CHỈ SỐ TRÊN TẬP ĐÃ XÁC THỰC ĐÚNG NGUỒN ({vm['verified_samples']} BÀI):")
            print(f"   ⭐ Verified Accuracy  : {vm['accuracy'] * 100:.2f}%")
            print(f"   🎯 Verified Precision : {vm['precision'] * 100:.2f}%")
            print(f"   📡 Verified Recall    : {vm['recall'] * 100:.2f}%")
            print(f"   ⚖️ Verified F1-Score  : {vm['f1_score'] * 100:.2f}%")
        print("!" * 75)

    # In danh sách bài sai (Error Analysis)
    wrong_items = [r for r in results if r["is_correct"] == "SAI"]
    if wrong_items:
        print("\n🔍 PHÂN TÍCH CÁC TRƯỜNG HỢP DỰ ĐOÁN SAI:")
        for idx, w in enumerate(wrong_items, 1):
            src_note = " [⚠️ SAI LỆCH NGUỒN]" if not w.get("is_verified_source") else ""
            print(f"  {idx}. '{w['song_title']}' - {w['artist']}{src_note}")
            print(
                f"     Thực tế: {w['ground_truth']} | Model AI: {w.get('ai_model_verdict', 'N/A')} | Tổng hợp: {w['prediction']} ({w['classification_type']})"
            )
            print(f"     Lý do model đưa ra: {w['reasons']}")
    else:
        print("\n🎉 Tuyệt vời! Không có bài nào bị dự đoán sai trong tập mẫu kiểm tra!")
    print("=" * 75)


def export_audit_excel(metrics: dict[str, Any], results: list[dict[str, Any]]) -> str:
    """Xuất file Excel đối soát chi tiết kết quả benchmark kèm Sheet 3 Sự cố & Nạp bù."""
    out_dir = Path("data/outputs")
    out_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_path = out_dir / f"benchmark_audit_{timestamp}.xlsx"

    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        # Sheet 1: Chi tiết
        df_details = pd.DataFrame(results)
        df_details.to_excel(writer, sheet_name="Audit Details", index=False)

        # Tính toán độ chính xác riêng của mô hình AI thuần túy
        ai_model_correct = sum(1 for r in results if r.get("ai_model_is_correct") == "ĐÚNG")
        ai_model_total = sum(1 for r in results if r.get("ai_model_is_correct") in ("ĐÚNG", "SAI"))
        ai_model_acc_str = (
            f"{(ai_model_correct / ai_model_total * 100):.2f}%" if ai_model_total > 0 else "N/A"
        )

        # Sheet 2: Tóm tắt
        summary_rows = [
            {"Chỉ số": "Tổng số bài mẫu thành công", "Giá trị": metrics["total_samples"]},
            {"Chỉ số": "Số bài thành công", "Giá trị": metrics["evaluated_successfully"]},
            {
                "Chỉ số": "Số bài gặp sự cố (Đã nạp bù)",
                "Giá trị": metrics.get("replaced_errors_count", metrics["errors"]),
            },
            {
                "Chỉ số": "Số bài sai lệch nguồn âm thanh (Unverified)",
                "Giá trị": metrics.get("unverified_matches", 0),
            },
            {
                "Chỉ số": "Độ chính xác Thuần AI Model (Raw Gemini)",
                "Giá trị": ai_model_acc_str,
            },
            {
                "Chỉ số": "Độ chính xác Tổng hợp (Consensus Final)",
                "Giá trị": f"{metrics['accuracy'] * 100:.2f}%",
            },
            {"Chỉ số": "True Positive (Đoán đúng AI)", "Giá trị": metrics["TP"]},
            {"Chỉ số": "True Negative (Đoán đúng Người)", "Giá trị": metrics["TN"]},
            {"Chỉ số": "False Positive (Vu oan Người)", "Giá trị": metrics["FP"]},
            {"Chỉ số": "False Negative (Bỏ sót AI)", "Giá trị": metrics["FN"]},
            {
                "Chỉ số": "Accuracy (Độ chính xác AI)",
                "Giá trị": f"{metrics['accuracy'] * 100:.2f}%",
            },
            {
                "Chỉ số": "Precision (Độ tin cậy AI)",
                "Giá trị": f"{metrics['precision'] * 100:.2f}%",
            },
            {"Chỉ số": "Recall (Độ nhạy bắt AI)", "Giá trị": f"{metrics['recall'] * 100:.2f}%"},
            {
                "Chỉ số": "Specificity (Bảo vệ nhạc người)",
                "Giá trị": f"{metrics['specificity'] * 100:.2f}%",
            },
            {"Chỉ số": "F1-Score (Phân biệt AI)", "Giá trị": f"{metrics['f1_score'] * 100:.2f}%"},
        ]

        wt = metrics.get("work_type_metrics")
        if wt and wt.get("total_evaluated", 0) > 0:
            summary_rows.extend(
                [
                    {
                        "Chỉ số": "──────── PHÂN BIỆT TÁC PHẨM ────────",
                        "Giá trị": "───────────────────",
                    },
                    {
                        "Chỉ số": "Tổng số bài đánh giá Tác phẩm",
                        "Giá trị": wt["total_evaluated"],
                    },
                    {
                        "Chỉ số": "Độ chính xác Tác phẩm (Work Type Accuracy)",
                        "Giá trị": f"{wt['accuracy'] * 100:.2f}%",
                    },
                ]
            )
            per_class = wt.get("per_class", {})
            for lbl, c_met in per_class.items():
                summary_rows.append(
                    {
                        "Chỉ số": f"Nhãn {lbl} (Precision / Recall / F1)",
                        "Giá trị": (
                            f"{c_met['precision'] * 100:.1f}% / "
                            f"{c_met['recall'] * 100:.1f}% / "
                            f"{c_met['f1_score'] * 100:.1f}%"
                        ),
                    }
                )
            if metrics.get("work_type_cm_image_path"):
                summary_rows.append(
                    {
                        "Chỉ số": "File ảnh Ma Trận Nhầm Lẫn Tác Phẩm",
                        "Giá trị": metrics["work_type_cm_image_path"],
                    }
                )

        df_summary = pd.DataFrame(summary_rows)
        df_summary.to_excel(writer, sheet_name="Executive Summary", index=False)

        # Sheet 3: Sự cố & Nạp bù (Error Audit & Replacement)
        error_recs = metrics.get("error_records", [])
        if error_recs:
            df_err = pd.DataFrame(error_recs)
            df_err.to_excel(writer, sheet_name="Sự cố & Nạp bù", index=False)

    return str(file_path)


def main():
    parser = argparse.ArgumentParser(
        description="Chấm điểm độ chính xác AI Music Detector với testcase.xlsx"
    )
    parser.add_argument(
        "--file", default="testcase.xlsx", help="Đường dẫn file excel (mặc định: testcase.xlsx)"
    )
    parser.add_argument(
        "--sheet", default="Media_W38", help="Tên sheet chứa dữ liệu (mặc định: Media_W38)"
    )
    parser.add_argument(
        "--sample", type=int, default=20, help="Số lượng bài mẫu muốn test (mặc định: 20)"
    )
    parser.add_argument(
        "--balanced", action="store_true", default=True, help="Lấy mẫu cân bằng 50/50 AI và Người"
    )
    parser.add_argument("--all", action="store_true", help="Chạy toàn bộ dataset (bỏ qua --sample)")
    parser.add_argument(
        "--model", type=str, default=None, help="Mô hình Gemini override (tùy chọn)"
    )
    parser.add_argument(
        "--force-rescan",
        action="store_true",
        default=False,
        help="Bắt buộc quét mới, không dùng cache cũ",
    )
    parser.add_argument(
        "--check-work-type",
        action="store_true",
        help="Kiểm tra và đối soát độ chính xác phân biệt Original / Cover / Remix với Ground Truth Song Type",
    )

    args = parser.parse_args()

    if args.check_work_type:
        from check_work_type_accuracy import check_work_type_accuracy

        check_work_type_accuracy(testcase_path=args.file)
        return

    print(f"📖 Đang đọc dữ liệu từ '{args.file}' [Sheet: '{args.sheet}']...")
    dataset = load_testcase_dataset(file_path=args.file, sheet_name=args.sheet)

    total_ai = sum(1 for d in dataset if d["gt_is_ai"])
    total_human = sum(1 for d in dataset if not d["gt_is_ai"])
    print(f"✅ Đã tải {len(dataset)} bài hát (AI: {total_ai} bài | Con người: {total_human} bài)")

    if args.all:
        samples = dataset
        print(f"⚠️ Chế độ FULL: Sẽ thẩm định toàn bộ {len(samples)} bài hát.")
    else:
        samples = select_samples(dataset, sample_size=args.sample, balanced=args.balanced)
        sample_ai = sum(1 for s in samples if s["gt_is_ai"])
        sample_human = sum(1 for s in samples if not s["gt_is_ai"])
        print(
            f"🎯 Lấy mẫu thẩm định: {len(samples)} bài (AI: {sample_ai} bài | Con người: {sample_human} bài)"
        )

    benchmark_data = run_benchmark(
        samples, dataset=dataset, model_override=args.model, force_rescan=args.force_rescan
    )
    metrics = benchmark_data["metrics"]
    results = benchmark_data["results"]

    print_summary_report(metrics, results)
    audit_file = export_audit_excel(metrics, results)
    print(f"\n📁 Đã xuất báo cáo đối soát chi tiết tại: {audit_file}\n")


if __name__ == "__main__":
    main()
