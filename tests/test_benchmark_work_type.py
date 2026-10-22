from __future__ import annotations

import unittest
from pathlib import Path

from benchmark_evaluator import (
    calculate_work_type_confusion_matrix,
    map_song_type_to_work_type,
    plot_work_type_confusion_matrix,
)


class TestBenchmarkWorkType(unittest.TestCase):
    """Kiểm thử tính năng đánh giá & vẽ ma trận nhầm lẫn cho Work Type (Original vs Cover vs Remix)."""

    def test_map_song_type_to_work_type(self) -> None:
        """Kiểm tra hàm ánh xạ Song Type từ testcase sang nhãn chuẩn."""
        self.assertEqual(map_song_type_to_work_type("Original"), "ORIGINAL")
        self.assertEqual(map_song_type_to_work_type("original"), "ORIGINAL")
        self.assertEqual(map_song_type_to_work_type("Live"), "ORIGINAL")
        self.assertEqual(map_song_type_to_work_type("Cover"), "COVER")
        self.assertEqual(map_song_type_to_work_type("Remake/Remix"), "REMIX")
        self.assertEqual(map_song_type_to_work_type("remix"), "REMIX")
        self.assertEqual(map_song_type_to_work_type(""), "UNCERTAIN")
        self.assertEqual(map_song_type_to_work_type(None), "UNCERTAIN")

    def test_calculate_work_type_confusion_matrix(self) -> None:
        """Kiểm tra tính toán ma trận nhầm lẫn đa lớp và chỉ số per-class."""
        y_true = ["ORIGINAL", "ORIGINAL", "COVER", "COVER", "REMIX", "REMIX"]
        y_pred = ["ORIGINAL", "REMIX", "COVER", "COVER", "REMIX", "ORIGINAL"]

        metrics = calculate_work_type_confusion_matrix(y_true, y_pred)

        self.assertEqual(metrics["total_evaluated"], 6)
        self.assertEqual(metrics["correct"], 4)
        self.assertAlmostEqual(metrics["accuracy"], 4 / 6, places=3)

        # Kiểm tra lớp COVER: 2/2 đoán đúng (Precision = 1.0, Recall = 1.0)
        cover_metrics = metrics["per_class"]["COVER"]
        self.assertEqual(cover_metrics["tp"], 2)
        self.assertEqual(cover_metrics["fn"], 0)
        self.assertEqual(cover_metrics["fp"], 0)
        self.assertEqual(cover_metrics["precision"], 1.0)
        self.assertEqual(cover_metrics["recall"], 1.0)

        # Kiểm tra ma trận
        mat = metrics["matrix"]
        self.assertEqual(mat["ORIGINAL"]["ORIGINAL"], 1)
        self.assertEqual(mat["ORIGINAL"]["REMIX"], 1)
        self.assertEqual(mat["COVER"]["COVER"], 2)
        self.assertEqual(mat["REMIX"]["REMIX"], 1)
        self.assertEqual(mat["REMIX"]["ORIGINAL"], 1)

    def test_plot_work_type_confusion_matrix(self) -> None:
        """Kiểm tra hàm vẽ và xuất file ảnh ma trận nhầm lẫn Work Type."""
        y_true = ["ORIGINAL", "COVER", "REMIX", "ORIGINAL"]
        y_pred = ["ORIGINAL", "COVER", "REMIX", "COVER"]

        metrics = calculate_work_type_confusion_matrix(y_true, y_pred)
        out_path = Path("data/temp/test_work_type_cm.png")

        generated_file = plot_work_type_confusion_matrix(
            metrics, output_path=out_path, title="Test Work Type Confusion Matrix"
        )

        self.assertTrue(generated_file.exists())
        self.assertGreater(generated_file.stat().st_size, 5000)  # File PNG có dữ liệu > 5KB

    def test_export_audit_excel_with_work_type(self) -> None:
        """Kiểm tra xuất file Excel chứa đầy đủ các chỉ số và chi tiết phân biệt Tác phẩm."""
        import openpyxl

        from benchmark_evaluator import export_audit_excel

        metrics = {
            "total_samples": 4,
            "evaluated_successfully": 4,
            "errors": 0,
            "TP": 2,
            "TN": 2,
            "FP": 0,
            "FN": 0,
            "accuracy": 1.0,
            "precision": 1.0,
            "recall": 1.0,
            "specificity": 1.0,
            "f1_score": 1.0,
            "work_type_metrics": {
                "total_evaluated": 4,
                "correct": 3,
                "accuracy": 0.75,
                "per_class": {
                    "ORIGINAL": {
                        "total": 2,
                        "tp": 2,
                        "precision": 1.0,
                        "recall": 1.0,
                        "f1_score": 1.0,
                    },
                    "COVER": {
                        "total": 1,
                        "tp": 1,
                        "precision": 1.0,
                        "recall": 1.0,
                        "f1_score": 1.0,
                    },
                    "REMIX": {
                        "total": 1,
                        "tp": 0,
                        "precision": 0.0,
                        "recall": 0.0,
                        "f1_score": 0.0,
                    },
                },
            },
            "work_type_cm_image_path": "data/outputs/work_type_cm_dummy.png",
        }
        results = [
            {
                "media_id": "row_1",
                "song_title": "Em của ngày hôm qua",
                "artist": "Sơn Tùng M-TP",
                "ground_truth": "HUMAN",
                "gt_work_type": "ORIGINAL",
                "pred_work_type": "ORIGINAL",
                "work_is_correct": "ĐÚNG",
                "work_label": "Bản gốc (Original)",
                "original_artist": "Sơn Tùng M-TP",
                "original_song_title": "Em của ngày hôm qua",
                "prediction": "HUMAN",
                "is_correct": "ĐÚNG",
            }
        ]

        excel_path = export_audit_excel(metrics, results)
        self.assertTrue(Path(excel_path).exists())

        wb = openpyxl.load_workbook(excel_path, data_only=True)
        self.assertIn("Executive Summary", wb.sheetnames)
        self.assertIn("Audit Details", wb.sheetnames)

        summary_sheet = wb["Executive Summary"]
        sheet_text = [str(row[0]) for row in summary_sheet.iter_rows(values_only=True)]
        self.assertTrue(any("Độ chính xác Tác phẩm" in t for t in sheet_text))


if __name__ == "__main__":
    unittest.main()
